"""Results-export tests: JSON and Excel (.xlsx).

These verify that the exporters run, produce a readable file, and preserve the
key numbers. The Excel test is skipped automatically if openpyxl is unavailable.
"""
import json
import os
import tempfile
import unittest

from context import Structure2D  # noqa: F401  (adds src/ to sys.path)
from xdfem2d import save_json, save_excel


def _beam():
    s = Structure2D()
    s.add_material("M", elastic_modulus=30e6, unit_weight=25.0)
    s.add_section("S", "M", b=0.3, h=0.6)
    s.add_node("N1", 0.0, 0.0)
    s.add_node("N2", 5.0, 0.0)
    s.add_bar_element("E1", "N1", "N2", "S")
    s.add_support("PIN", ux=True, uy=True)
    s.add_support("ROLLER", ux=False, uy=True)
    s.assign_support("N1", "PIN")
    s.assign_support("N2", "ROLLER")
    s.add_load_case("LC")
    s.add_distributed_load("E1", "LC", fye=-10.0, fyd=-10.0)
    # An analysis case is what the result tables report (load cases are inputs).
    s.add_analysis_case("AC", "Linear", {"LC": 1.0})
    s.add_load_combination("C", {"LC": 1.4}, combo_type="LinearSum")
    return s


def _solved_beam():
    return _beam().calculate()


class TestSaveJson(unittest.TestCase):
    """save_json must serialise NumPy arrays/scalars (regression guard)."""

    def test_json_roundtrip(self):
        r = _solved_beam()
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "results.json")
            save_json(r, path)
            self.assertGreater(os.path.getsize(path), 0)
            with open(path, encoding="utf-8") as f:
                back = json.load(f)
        # Key numbers preserved through the round-trip.
        self.assertAlmostEqual(back["reactions"]["LC"]["N1"][1], 25.0, places=6)
        self.assertAlmostEqual(back["reactions"]["LC"]["N2"][1], 25.0, places=6)
        # Diagram arrays (NumPy) survive as JSON lists.
        self.assertIn("element_distribution", back)
        self.assertIsInstance(back["element_distribution"]["LC"]["E1"]["M"], list)


class TestSaveExcel(unittest.TestCase):
    def setUp(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")

    def test_xlsx_has_expected_sheets(self):
        from openpyxl import load_workbook
        r = _solved_beam()
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "results.xlsx")
            save_excel(r, path)
            self.assertGreater(os.path.getsize(path), 0)
            wb = load_workbook(path)
        for sheet in ("Displacements", "Reactions", "Element forces"):
            self.assertIn(sheet, wb.sheetnames)

    def test_xlsx_reaction_value_written(self):
        from openpyxl import load_workbook
        r = _solved_beam()
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "results.xlsx")
            save_excel(r, path)
            ws = load_workbook(path)["Reactions"]
            rows = [tuple(c.value for c in row) for row in ws.iter_rows()]
        # Header: [Case type, Case, Node, Rx, Ry, Mz] → Ry is column index 4.
        ry_values = [round(row[4], 3) for row in rows[1:]
                     if isinstance(row[4], (int, float))]
        self.assertIn(25.0, ry_values)

    def test_save_structure_excel(self):
        from openpyxl import load_workbook
        from xdfem2d import save_structure_excel
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "model.xlsx")
            save_structure_excel(_beam(), path)
            wb = load_workbook(path)
        for sheet in ("Nodes", "Beams", "Supports"):   # topology names
            self.assertIn(sheet, wb.sheetnames)

    def test_passing_the_struc_gives_plate_headers(self):
        """A plate model's results save with plate headers (w, Rz) when the
        model is passed, without the caller spelling out the domain."""
        from openpyxl import load_workbook
        from xdfem2d import model_json as MJ
        s = MJ.load(MJ.template("slab", as_text=False))
        r = s.calculate()
        with tempfile.TemporaryDirectory() as d:
            # Default (no struc, no domain) is plane — the old behaviour.
            plane_path = os.path.join(d, "plane.xlsx")
            save_excel(r, plane_path)
            plane_disp = load_workbook(plane_path)["Displacements"]
            plane_hdr = [c.value for c in next(plane_disp.iter_rows())]
            self.assertTrue(any("ux" in str(h) for h in plane_hdr))

            # Passing the model reads the domain and writes plate headers.
            plate_path = os.path.join(d, "plate.xlsx")
            save_excel(r, plate_path, struc=s)
            wb = load_workbook(plate_path)
            disp_hdr = [c.value for c in next(wb["Displacements"].iter_rows())]
            reac_hdr = [c.value for c in next(wb["Reactions"].iter_rows())]
        self.assertTrue(any("w [m]" == str(h) for h in disp_hdr), disp_hdr)
        self.assertTrue(any("Rz" in str(h) for h in reac_hdr), reac_hdr)


if __name__ == "__main__":
    unittest.main(verbosity=2)
