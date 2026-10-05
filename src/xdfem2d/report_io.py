"""
xdfem2D report / export builders (GUI-free).

Turn a solved results dict (and the model) into tables, plain-text report lines,
and Excel/Word/PDF documents. Lives in the API so it can be used from scripts
and tests; the GUI imports these helpers.
"""
from __future__ import annotations

import numpy as np


def _float(text, default=0.0) -> float:
    """Parse a value to float, returning *default* on failure."""
    try:
        return float(str(text).strip())
    except (ValueError, TypeError):
        return default


# ── Table number formatting ─────────────────────────────────────────────────
# Fixed decimal places per physical unit, keyed by the token inside a header's
# "[...]". A column's precision is read from its header so every value in it
# lines up; anything without a known unit falls back to a compact 3 sig-figs.
_UNIT_DECIMALS = {
    'm': 3, 'm²': 3, 'm³': 3, 'm⁴': 6, 'mm': 1,
    'rad': 5, '°': 1, '1/°c': 6, '°c': 1,
    'kn': 2, 'knm': 2, 'kn/m': 1, 'kn/m²': 1, 'knm/m': 1, 'kn/m³': 0,
    't': 3, 't·m²': 3, 'hz': 3, 's': 3, 'rad/s': 3,
    '%': 1, 'cm²/m': 2, 'cm²': 2, 'mpa': 2,
}


def _header_decimals(header: str):
    """Decimal count for a column, from the unit token in its header (or None)."""
    import re
    m = re.search(r'\[([^\]]+)\]', str(header))
    if m:
        return _UNIT_DECIMALS.get(m.group(1).strip().lower())
    return None


def _fmt_num(v, decimals):
    """Format a float with fixed *decimals* (or compact 3 sig-figs if unknown)."""
    if decimals is None:
        return f"{v:.3g}"
    return f"{v:.{decimals}f}"


def _fmt_cell(v, decimals):
    """One cell as text: fixed-decimal floats, blank for None, str otherwise."""
    if v is None or v == "":
        return ""
    if isinstance(v, float):
        return _fmt_num(v, decimals)
    return str(v)


def _numeric_columns(header, rows):
    """Set of column indices whose every non-empty cell is a number — these are
    right-aligned (decimals line up); text columns stay left-aligned."""
    num = set(range(len(header)))
    for r in rows:
        for i, v in enumerate(r):
            if i in num and v not in (None, "") and not isinstance(v, (int, float)):
                num.discard(i)
    return num


def _sum_reactions(reac: dict):
    """Return (ΣRx, ΣRy, ΣMz) over a per-node reactions dict, or None."""
    if not isinstance(reac, dict):
        return None
    vals = [r for r in reac.values() if isinstance(r, list) and len(r) >= 3]
    if not vals:
        return None
    return (sum(r[0] for r in vals), sum(r[1] for r in vals), sum(r[2] for r in vals))


def _results_reaction_sum_rows(results: dict):
    """(kind, case, ΣRx, ΣRy, ΣMz) per combination / analysis case (load cases
    omitted). Envelopes are skipped (per-node extremes don't sum physically)."""
    rows = []
    for cid, cres in results.get('combinations', {}).items():
        reac = cres.get('reactions', {})
        if isinstance(reac, dict) and 'max' in reac:
            continue
        s = _sum_reactions(reac)
        if s:
            rows.append(("Combination", cid, *s))
    for ac_id, ac_res in results.get('analysis_cases', {}).items():
        if not isinstance(ac_res, dict) or 'error' in ac_res:
            continue
        s = _sum_reactions(ac_res.get('reactions', {}))
        if s:
            rows.append(("Analysis case", ac_id, *s))
    return rows


def _results_modal_rows(results: dict):
    """(case, mode, ω, f, T, Meff_X%, Meff_Y%, ΣX%, ΣY%) for every modal case."""
    rows = []
    for ac_id, ac_res in results.get('analysis_cases', {}).items():
        if not isinstance(ac_res, dict) or 'modal_info' not in ac_res:
            continue
        cum_x = cum_y = 0.0
        for mi in ac_res['modal_info']:
            cum_x += mi['meff_x_pct']
            cum_y += mi['meff_y_pct']
            rows.append((ac_id, mi['mode'], mi['omega'], mi['frequency'],
                         mi['period'], mi['meff_x_pct'], mi['meff_y_pct'],
                         cum_x, cum_y))
    return rows


def _pdf_text_pages(pdf, lines: list, title: str, stamp: str,
                    section_break: bool = True):
    """Append paginated monospaced A4 text pages for `lines` to an open PdfPages."""
    from matplotlib.figure import Figure
    fig_h = 11.69
    fontsize = 7.5
    body_top = 0.945
    body_bot = 0.04
    line_frac = (fontsize / 72.0) / fig_h * 1.42
    per_page = max(1, int((body_top - body_bot) / line_frac))

    BAR = '=' * 50
    if section_break:
        n = len(lines)
        sections, cur, i = [], [], 0
        while i < n:
            if lines[i] == BAR and i + 2 < n and lines[i + 2] == BAR:
                if cur:
                    sections.append(cur)
                cur = [lines[i]]
            else:
                cur.append(lines[i])
            i += 1
        if cur:
            sections.append(cur)
        if not sections:
            sections = [lines]
    else:
        sections = [lines]

    for sec in sections:
        for p in range(0, len(sec), per_page):
            chunk = sec[p:p + per_page]
            fig = Figure(figsize=(8.27, fig_h))
            ax = fig.add_axes([0, 0, 1, 1]); ax.axis('off')
            ax.text(0.06, 0.975, title, va='top', ha='left', fontsize=12,
                    fontweight='bold')
            ax.text(0.94, 0.975, stamp, va='top', ha='right',
                    fontsize=8, color='gray')
            ax.text(0.06, body_top, "\n".join(chunk), va='top', ha='left',
                    family='monospace', fontsize=fontsize, linespacing=1.42)
            pdf.savefig(fig)


def _export_results_to_pdf(lines: list, path: str, title: str = "xdfem2D — Results"):
    """Render the plain-text results report to a paginated A4 PDF (monospace)."""
    from matplotlib.backends.backend_pdf import PdfPages
    import datetime as _dt
    stamp = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    with PdfPages(path) as pdf:
        _pdf_text_pages(pdf, lines, title, stamp, section_break=True)


def _build_model_lines(struc, include_project=True, include_model=True) -> list:
    """Plain-text dump of the model (input) data, one table per section."""
    g = getattr
    L = []

    def hdr(t):
        L.extend(['=' * 50, f"  {t}", '=' * 50])

    def row(s):
        L.append(f"    {s}")

    info = getattr(struc, 'project_info', {}) or {}
    if include_project and any(str(v).strip() for v in info.values()):
        hdr("PROJECT")
        for k, v in info.items():
            if str(v).strip():
                row(f"{k}: {v}")
        L.append("")
    if not include_model:
        return L
    hdr("MATERIALS")
    for m in struc.materials.values():
        row(f"{m.name}: E={g(m,'elastic_modulus',''):g}  "
            f"γ={g(m,'unit_weight',''):g}  α={g(m,'alpha',0):g}")
    L.append("")
    hdr("SECTIONS")
    for s in struc.sections.values():
        row(f"{s.name}: material={g(s,'material_name','')}  "
            f"b={g(s,'b','')}  h={g(s,'h','')}")
    L.append("")
    hdr("NODES")
    for nid, n in struc.nodes.items():
        row(f"{nid}: ({n.x:g}, {n.y:g})")
    L.append("")
    hdr("ELEMENTS")
    for e in struc.bar_elements:
        row(f"{e.id}: {e.node_i} → {e.node_j}  section={e.section_name}  "
            f"hinge_i={bool(g(e,'hinge_i',False))}  hinge_j={bool(g(e,'hinge_j',False))}")
    L.append("")
    hdr("SUPPORTS")
    for a in struc.support_assignments:
        sp = struc.supports.get(a.support_name)
        row(f"{a.node_id}: {a.support_name}  "
            f"UX={bool(g(sp,'ux',False))} UY={bool(g(sp,'uy',False))} TZ={bool(g(sp,'tz',False))}")
    L.append("")
    if struc.node_springs:
        hdr("NODE SPRINGS")
        _pl = g(struc, 'domain', 'plane') == 'plate'
        _k = ('Kz', 'Kθx', 'Kθy') if _pl else ('Kx', 'Ky', 'Kt')
        for nid, sp in struc.node_springs.items():
            row(f"{nid}: {_k[0]}={g(sp,'kx',0):g}  {_k[1]}={g(sp,'ky',0):g}  "
                f"{_k[2]}={g(sp,'kt',0):g}")
        L.append("")
    if struc.element_springs:
        hdr("ELEMENT SPRINGS")
        for eid, es in struc.element_springs.items():
            if g(struc, 'domain', 'plane') == 'plate':   # only the w spring acts
                row(f"{eid}: Kz={g(es,'ky',0):g}")
            else:
                row(f"{eid}: Kx={g(es,'kx',0):g}  Ky={g(es,'ky',0):g}  "
                    f"axes={g(es,'coord_sys','global')}")
        L.append("")
    hdr("LOAD CASES")
    for lc in struc.load_cases:
        row(f"{lc.id}: action={str(g(lc,'action_type',''))}  "
            f"self-weight={g(lc,'self_weight_factor',0):g}")
    L.append("")
    if struc.point_loads:
        hdr("POINT LOADS")
        _pl = g(struc, 'domain', 'plane') == 'plate'
        _n = ('Fz', 'Mx', 'My') if _pl else ('Fx', 'Fy', 'Mz')
        for p in struc.point_loads:
            row(f"{p.node_id} [{p.load_case_id}]: {_n[0]}={p.fx:g}  "
                f"{_n[1]}={p.fy:g}  {_n[2]}={p.mz:g}")
        L.append("")
    if struc.distributed_loads:
        hdr("DISTRIBUTED LOADS")
        for d in struc.distributed_loads:
            if g(struc, 'domain', 'plane') == 'plate':   # transverse qz only
                row(f"{d.element_id} [{d.load_case_id}]: qz={d.fye:g}→{d.fyd:g}")
            else:
                row(f"{d.element_id} [{d.load_case_id}]: Fy={d.fye:g}→{d.fyd:g}  "
                    f"Fx={d.fxe:g}→{d.fxd:g}  ({d.coord_sys})")
        L.append("")
    if getattr(struc, 'element_point_loads', None):
        hdr("ELEMENT POINT LOADS")
        for e in struc.element_point_loads:
            if g(struc, 'domain', 'plane') == 'plate':   # Fz + bending moment
                row(f"{e.element_id} [{e.load_case_id}]: a={e.a:g}m  "
                    f"Fz={e.fy:g}  M={e.mz:g}")
            else:
                row(f"{e.element_id} [{e.load_case_id}]: a={e.a:g}m  "
                    f"Fx={e.fx:g}  Fy={e.fy:g}  Mz={e.mz:g}  ({e.coord_sys})")
        L.append("")
    if struc.temperature_loads:
        hdr("TEMPERATURE LOADS")
        for t in struc.temperature_loads:
            row(f"{t.element_id} [{t.load_case_id}]: ΔT={t.delta_t_uniform:g}  "
                f"gradient={t.delta_t_gradient:g}")
        L.append("")
    if struc.support_settlements:
        hdr("SETTLEMENTS")
        for s in struc.support_settlements:
            row(f"{s.node_id} [{s.load_case_id}]: UX={s.ux:g}  UY={s.uy:g}  TZ={s.tz:g}")
        L.append("")
    if struc.nodal_masses:
        hdr("NODAL MASSES")
        for m in struc.nodal_masses:
            row(f"{m.node_id} [{m.mass_case_id}]: Mx={m.mx:g}  My={m.my:g}  Mtz={m.mtz:g}")
        L.append("")
    hdr("ANALYSIS CASES")
    for a in struc.analysis_cases:
        coef = "  ".join(f"{k}×{v:g}" for k, v in g(a, 'coefficients', {}).items())
        row(f"{a.id}: {a.analysis_type}   {coef}")
    L.append("")
    if struc.load_combinations:
        hdr("COMBINATIONS")
        for c in struc.load_combinations:
            coef = "  ".join(f"{k}×{v:g}" for k, v in g(c, 'coefficients', {}).items())
            row(f"{c.id}: {g(c,'combo_type','')}   {coef}")
        L.append("")
    return L


def _model_tables(struc, include_project=True, include_model=True):
    """[(title, header, rows)] for the model (input) data — one table each."""
    g = getattr
    T = []
    if include_project:
        T.append(("Project", ["Field", "Value"],
                  [(k, v) for k, v in (g(struc, 'project_info', {}) or {}).items()]))
    if not include_model:
        return [t for t in T if t[2]]
    def _enum_val(v):
        return getattr(v, 'value', v) if v is not None else ''

    def _design_summary(m):
        design = g(m, 'design', None) or {}
        return "  ".join(f"{k}={v}" for k, v in design.items() if v not in (None, ""))

    # Topology vocabulary by domain: a line element is a Beam (plane) or Grid
    # member (plate); an area element is a Panel (plane) or Slab (plate). Restraint
    # and load components read (ux,uy,θz) in plane and (w,θx,θy) in plate.
    plate = g(struc, 'domain', 'plane') == 'plate'
    line_role = 'Grid' if plate else 'Beam'
    area_role = 'Slab' if plate else 'Panel'
    line_pl = line_role if plate else line_role + 's'     # Grid / Beams
    area_pl = area_role + 's'                              # Slabs / Panels
    d0, d1, d2 = ('W', 'θx', 'θy') if plate else ('UX', 'UY', 'TZ')
    pl = (['Fz [kN]', 'Mx [kNm]', 'My [kNm]'] if plate
          else ['Fx [kN]', 'Fy [kN]', 'Mz [kNm]'])
    stl = (['W [m]', 'θx [rad]', 'θy [rad]'] if plate
           else ['UX [m]', 'UY [m]', 'TZ [rad]'])

    T.append(("Materials",
              ["Name", "Type", "E [kN/m²]", "Poisson ν", "γ [kN/m³]", "α [1/°C]",
               "Unit mass", "Design"],
              [(m.name, _enum_val(g(m, 'material_type', '')),
                g(m, 'elastic_modulus', ''), g(m, 'poisson', 0.2), g(m, 'unit_weight', ''),
                g(m, 'alpha', ''), g(m, 'unit_mass', '') or "", _design_summary(m))
               for m in struc.materials.values()]))
    T.append(("Sections",
              ["Name", "Type", "Shape", "Material", "b [m]", "h [m]",
               "tw [m]", "tf [m]", "Area [m²]", "I [m⁴]", "Profile",
               "Cover [m]", "Stirrup α [°]", "Bar Ø [mm]"],
              [(s.name, _enum_val(g(s, 'section_type', '')),
                _enum_val(g(s, 'shape', '')), g(s, 'material_name', ''),
                g(s, 'b', ''), g(s, 'h', ''), g(s, 'tw', 0.0), g(s, 'tf', 0.0),
                g(s, 'area', ''), g(s, 'inertia', ''), g(s, 'profile_name', '') or "",
                g(s, 'rc_cover', ''), g(s, 'rc_alpha_s', ''), g(s, 'rc_bar_phi', 16.0))
               for s in struc.sections.values()]))
    T.append((f"{area_role} sections",
              ["Name", "Material", "Thickness [m]", "Model", "Formulation",
               "Cover [m]", "Stirrup α [°]", "Bar Ø [mm]", "Min. shear reinf."],
              [(ts.name, g(ts, 'material_name', ''), g(ts, 'thickness', ''),
                "plane strain" if g(ts, 'plane_strain', False) else "plane stress",
                g(ts, 'formulation', 'CST'), g(ts, 'rc_cover', ''),
                g(ts, 'rc_alpha_s', ''), g(ts, 'rc_bar_phi', 16.0),
                bool(g(ts, 'rc_shear_min', True)))
               for ts in struc.tri_sections.values()]))
    # Quad sections: a distinct collection from tri_sections (QuadSection is
    # its own dataclass -- see models.py), so a model that uses quad plate/
    # membrane elements (DKT4/MITC4/Q4/QM6) needs its own table here, same
    # columns as the triangle one.
    T.append((f"{area_role} sections (quad)",
              ["Name", "Material", "Thickness [m]", "Model", "Formulation",
               "Cover [m]", "Stirrup α [°]", "Bar Ø [mm]", "Min. shear reinf."],
              [(qs.name, g(qs, 'material_name', ''), g(qs, 'thickness', ''),
                "plane strain" if g(qs, 'plane_strain', False) else "plane stress",
                g(qs, 'formulation', 'MITC4'), g(qs, 'rc_cover', ''),
                g(qs, 'rc_alpha_s', ''), g(qs, 'rc_bar_phi', 16.0),
                bool(g(qs, 'rc_shear_min', True)))
               for qs in getattr(struc, 'quad_sections', {}).values()]))
    T.append(("Nodes", ["ID", "X [m]", "Y [m]"],
              [(nid, n.x, n.y) for nid, n in struc.nodes.items()]))
    T.append((line_pl, ["ID", "Node i", "Node j", "Section", "Hinge i", "Hinge j"],
              [(e.id, e.node_i, e.node_j, e.section_name,
                bool(g(e, 'hinge_i', False)), bool(g(e, 'hinge_j', False)))
               for e in struc.bar_elements]))
    T.append((area_pl, ["ID", "Node i", "Node j", "Node k", f"{area_role} section"],
              [(t.id, t.node_i, t.node_j, t.node_k, t.section_name)
               for t in getattr(struc, 'tri_elements', [])]))
    # Quad elements: a distinct collection from tri_elements (see
    # structure.py) -- was entirely missing from this export before (results
    # tables already merge quad rows into the triangle-stress tables, but
    # the model tables never listed the quads' own connectivity/section at
    # all). Own table since a quad has a 4th node the triangle table has no
    # column for.
    T.append((f"{area_pl} (quad)",
              ["ID", "Node i", "Node j", "Node k", "Node l", f"{area_role} section"],
              [(q.id, q.node_i, q.node_j, q.node_k, q.node_l, q.section_name)
               for q in getattr(struc, 'quad_elements', [])]))
    sup = []
    for a in struc.support_assignments:
        sp = struc.supports.get(a.support_name)
        sup.append((a.node_id, a.support_name,
                    bool(g(sp, 'ux', False)) if sp else "",
                    bool(g(sp, 'uy', False)) if sp else "",
                    bool(g(sp, 'tz', False)) if sp else ""))
    T.append(("Supports", ["Node", "Support", d0, d1, d2], sup))
    T.append(("Node springs",
              (["Node", "Kz [kN/m]", "Kθx [kNm/rad]", "Kθy [kNm/rad]"] if plate
               else ["Node", "Kx [kN/m]", "Ky [kN/m]", "Kt [kNm/rad]"]),
              [(nid, g(sp, 'kx', ''), g(sp, 'ky', ''), g(sp, 'kt', ''))
               for nid, sp in struc.node_springs.items()]))
    if plate:       # only the transverse (w) spring acts on a grillage bar
        T.append(("Element springs", ["Element", "Kz [kN/m²]"],
                  [(eid, g(sp, 'ky', '')) for eid, sp in struc.element_springs.items()]))
    else:
        T.append(("Element springs", ["Element", "Kx [kN/m²]", "Ky [kN/m²]", "Axes"],
                  [(eid, g(sp, 'kx', ''), g(sp, 'ky', ''), g(sp, 'coord_sys', 'global'))
                   for eid, sp in struc.element_springs.items()]))
    T.append(("Load cases", ["ID", "Action type", "Self-weight factor"],
              [(lc.id, str(g(lc, 'action_type', '')), g(lc, 'self_weight_factor', 0.0))
               for lc in struc.load_cases]))
    T.append(("Point loads", ["Node", "Load case", pl[0], pl[1], pl[2]],
              [(p.node_id, p.load_case_id, p.fx, p.fy, p.mz) for p in struc.point_loads]))
    if plate:       # a grillage bar: transverse qz, Fz and a bending moment only
        T.append(("Distributed loads",
                  ["Element", "Load case", "qz start", "qz end"],
                  [(d.element_id, d.load_case_id, d.fye, d.fyd)
                   for d in struc.distributed_loads]))
        T.append(("Element point loads",
                  ["Element", "Load case", "a [m]", "Fz [kN]", "M [kNm]"],
                  [(e.element_id, e.load_case_id, e.a, e.fy, e.mz)
                   for e in getattr(struc, 'element_point_loads', [])]))
    else:
        T.append(("Distributed loads",
                  ["Element", "Load case", "Fy start", "Fy end", "Fx start", "Fx end", "Coord"],
                  [(d.element_id, d.load_case_id, d.fye, d.fyd, d.fxe, d.fxd, d.coord_sys)
                   for d in struc.distributed_loads]))
        T.append(("Element point loads",
                  ["Element", "Load case", "a [m]", "Fx [kN]", "Fy [kN]", "Mz [kNm]", "Coord"],
                  [(e.element_id, e.load_case_id, e.a, e.fx, e.fy, e.mz, e.coord_sys)
                   for e in getattr(struc, 'element_point_loads', [])]))
    edge = [("tri", e) for e in getattr(struc, 'tri_edge_loads', [])]
    edge += [("surface", e) for e in getattr(struc, 'surface_edge_loads', [])]
    T.append(("Edge loads",
              ["ID", "On", "Owner", "Edge", "Load case", "Fx [kN/m]", "Fy [kN/m]"],
              [(e.id, kind, (e.tri_id if kind == "tri" else e.object_id),
                f"{e.node_a}–{e.node_b}", e.load_case_id, e.fx, e.fy)
               for kind, e in edge]))
    T.append(("Temperature loads",
              ["Element", "Load case", "ΔT uniform [°C]", "ΔT gradient [°C]"],
              [(t.element_id, t.load_case_id, t.delta_t_uniform, t.delta_t_gradient)
               for t in struc.temperature_loads]))
    T.append(("Settlements", ["Node", "Load case", stl[0], stl[1], stl[2]],
              [(s.node_id, s.load_case_id, s.ux, s.uy, s.tz)
               for s in struc.support_settlements]))
    T.append(("Nodal masses", ["Node", "Mass case", "Mx [t]", "My [t]", "Mtz [t·m²]"],
              [(m.node_id, m.mass_case_id, m.mx, m.my, m.mtz) for m in struc.nodal_masses]))
    T.append(("Analysis cases",
              ["ID", "Type", "Load-case coefficients", "Modal case", "Spectrum", "Stored K"],
              [(a.id, a.analysis_type,
                "  ".join(f"{k}×{v:g}" for k, v in g(a, 'coefficients', {}).items()),
                g(a, 'modal_case_id', '') or "", g(a, 'spectrum_id', '') or "",
                g(a, 'stored_stiffness_id', '') or "")
               for a in struc.analysis_cases]))
    T.append(("Combinations", ["ID", "Type", "Analysis-case coefficients"],
              [(c.id, g(c, 'combo_type', ''),
                "  ".join(f"{k}×{v:g}" for k, v in g(c, 'coefficients', {}).items()))
               for c in struc.load_combinations]))
    # Quantity take-off: counts, lengths, areas, volumes, weights and masses,
    # by element type, material and section (objects expanded to their mesh).
    from .model_stats import model_stats_rows
    T.append(("Model info", ["Label", "Value", "Units"],
              [(lbl, "" if val is None else val, unit or "")
               for lbl, val, unit in model_stats_rows(struc)]))
    return [t for t in T if t[2]]


# Titles _results_tables can produce that are specifically EC2 design output
# (reinforcement, punching, crack width, the design partial factors) rather
# than raw FEM results (displacements, reactions, element forces, stresses,
# modal). The Design report panel (design_report_panel.py, Design ▸ Design
# report…) is now the canonical place for these -- it builds them properly,
# per governing combination, from the same design routines (build_reports /
# filter_reports / summary_rows), the way the equivalent design DRAWINGS
# were already pulled out of the Write Report dialog and into that panel's
# General tab (see ExportDialog's own comment on that). Write Report
# (MainWindow._export) excludes these titles for the same reason; kept here,
# not private, so that exclusion has one shared source of truth instead of
# a second hardcoded copy of this list in main_window.py.
DESIGN_RESULT_TITLES = {
    "Design parameters",
    "Reinforcement (EC2 design)",
    "Grillage reinforcement (EC2 bending + shear + torsion)",
    "Slab reinforcement (EC2 flexure)",
    "Shell reinforcement (EC2 design)",
    "Punching shear (EN 1992-1-1:2023)",
    "Crack width (SLS)",
}


def _results_tables(results: dict, domain: str = "plane", scope=None):
    """[(title, header, rows)] for the analysis results — one table each.

    ``domain`` ('plane' | 'plate') sets the vocabulary: a plane model reports
    nodal ux/uy/θz, reactions Rx/Ry/Mz and membrane triangle stresses; a plate
    model reports w/θx/θy, reactions Rz/Mx/My and slab forces (moments, shears,
    Wood–Armer) — the same storage, read in the domain's terms.

    ``scope`` — None (default, unfiltered — every existing caller/test keeps
    its current output) or a dict ``{"nodes": set, "bars": set, "areas": set}``
    (``"areas"`` covers both triangles and quads under one set, since the
    tables below key by either kind under the same column) — restricts each
    per-node/per-element table to the given ids, e.g. the GUI's active Scene
    (see MainWindow._export_scope). Model-wide summaries that carry no single
    node/element identity per row (Reaction sums, Modal, Design parameters)
    are never filtered, scope or not — there is nothing in a row to test.
    """
    plate = domain == "plate"
    cases = list(_iter_result_cases(results))
    scope = scope or {}
    scoped = bool(scope)
    sc_nodes = scope.get("nodes")
    sc_bars = scope.get("bars")
    sc_areas = scope.get("areas")
    T = []
    disp_hdr = (["Case type", "Case", "Node", "w [m]", "θx [rad]", "θy [rad]"]
                if plate else
                ["Case type", "Case", "Node", "ux [m]", "uy [m]", "rz [rad]"])
    reac_hdr = (["Case type", "Case", "Node", "Rz [kN]", "Mx [kNm]", "My [kNm]"]
                if plate else
                ["Case type", "Case", "Node", "Rx [kN]", "Ry [kN]", "Mz [kNm]"])
    rows = []
    for kind, case, disp, _r, _e in cases:
        for nid, d in (disp or {}).items():
            if sc_nodes is not None and nid not in sc_nodes:
                continue
            if isinstance(d, list) and len(d) >= 3:
                rows.append([kind, case, nid, d[0], d[1], d[2]])
    T.append(("Displacements", disp_hdr, rows))
    rows = []
    sum_label = "Σ (scene supports)" if scoped else "Σ (all supports)"
    for kind, case, _d, reac, _e in cases:
        had = False
        reac_in_scope = reac
        if sc_nodes is not None:
            reac_in_scope = {nid: r for nid, r in (reac or {}).items()
                             if nid in sc_nodes}
        for nid, r in (reac_in_scope or {}).items():
            if isinstance(r, list) and len(r) >= 3:
                rows.append([kind, case, nid, r[0], r[1], r[2]]); had = True
        # The sum is over whichever supports actually fed the rows above —
        # every one when unscoped, only the in-scope ones otherwise, so it
        # never silently claims to be "all supports" when it isn't.
        s = _sum_reactions(reac_in_scope)
        if had and s:
            rows.append([kind, case, sum_label, s[0], s[1], s[2]])
    T.append(("Reactions", reac_hdr, rows))
    rows = []
    for kind, case, _d, _r, ef in cases:
        for eid, v in (ef or {}).items():
            if sc_bars is not None and eid not in sc_bars:
                continue
            if isinstance(v, dict) and 'i' in v and 'j' in v:
                rows.append([kind, case, eid, "i", v['i'][0], v['i'][1], v['i'][2]])
                rows.append([kind, case, eid, "j", v['j'][0], v['j'][1], v['j'][2]])
    # Domain-aware header: a plate/grillage bar has no axial force -- the
    # first slot is torsion T instead (see steel_design.py/timber_design.py/
    # rc_design.py's own "Grillage: the 'N' slot is torsion" comments, and
    # the Grillage reinforcement table below, which already relabels this
    # as T_Ed). Same values, same column order -- only the label/unit for
    # that first column changes, matching how disp_hdr/reac_hdr above are
    # already domain-aware.
    ef_hdr = (["Case type", "Case", "Element", "End", "T [kNm]", "V [kN]", "M [kNm]"]
              if plate else
              ["Case type", "Case", "Element", "End", "N [kN]", "V [kN]", "M [kNm]"])
    T.append(("Element forces", ef_hdr, rows))

    # ── Constraint forces (multi-point constraints) ───────────────────
    # The force each constraint applies at the nodes it couples (method 2:
    # residual of the physical stiffness). Empty tables are dropped at the end,
    # so this appears only for models that have constraints.
    rows = []
    def _cf_emit(kind, case, mp):
        for nid, v in (mp or {}).items():
            if sc_nodes is not None and nid not in sc_nodes:
                continue
            if isinstance(v, list) and len(v) >= 3:
                rows.append([kind, case, nid, v[0], v[1], v[2]])
    def _cf_rows(container, kind):
        for case, cr in (container or {}).items():
            if not isinstance(cr, dict):
                continue
            cf = cr.get('constraint_forces')
            # Envelope / seismic combinations carry a max/min band; split it into
            # two labelled rows, exactly as the Reactions table does.
            if isinstance(cf, dict) and 'max' in cf:
                _cf_emit(kind, f"{case} (MAX)", cf.get('max'))
                _cf_emit(kind, f"{case} (MIN)", cf.get('min'))
            else:
                _cf_emit(kind, case, cf)
    _cf_rows(results.get('analysis_cases'), "Analysis")
    _cf_rows(results.get('combinations'), "Combination")
    T.append(("Constraint forces",
              ["Case type", "Case", "Node", "Fx [kN]", "Fy [kN]", "Mz [kNm]"],
              rows))

    rows = []
    for ac, acr in results.get('analysis_cases', {}).items():
        if not isinstance(acr, dict):
            continue
        for eid, d in acr.get('element_distribution', {}).items():
            if sc_bars is not None and eid not in sc_bars:
                continue
            try:
                rows.append([ac, eid, float(np.max(np.abs(d['N']))),
                             float(np.max(np.abs(d['V']))),
                             float(np.max(d['M'])), float(np.min(d['M']))])
            except Exception:
                pass
    T.append(("Span extremes", ["Case", "Element", "|N|max [kN]", "|V|max [kN]", "M+ [kNm]", "M- [kNm]"], rows))
    T.append(("Reaction sums", ["Case type", "Case", "ΣRx", "ΣRy", "ΣMz"],
              [list(r) for r in _results_reaction_sum_rows(results)]))
    T.append(("Modal", ["Case", "Mode", "ω [rad/s]", "f [Hz]", "T [s]",
                        "Meff_X %", "Meff_Y %", "Cum X %", "Cum Y %"],
              [list(r) for r in _results_modal_rows(results)]))

    # ── Triangle (CST + Allman) stresses — per analysis case / combination ──
    # 'formulation' is stamped onto each stress dict by tri_stresses/
    # allman_stresses themselves (see tri_elements.py / tri_elements_allman.py),
    # so this table doesn't need the Structure2D object — only `results`.
    rows = []
    if plate:
        # Slab forces: bending/twisting moments, principal moments, shears and
        # the Wood–Armer (W-A) design moments — the plate tri_stress dict keys.
        def _tri_rows(container, kind):
            for case, cr in (container or {}).items():
                if not isinstance(cr, dict):
                    continue
                for tid, d in (cr.get('tri_stress') or {}).items():
                    if sc_areas is not None and tid not in sc_areas:
                        continue
                    rows.append([kind, case, tid,
                                 d.get('mx'), d.get('my'), d.get('mxy'),
                                 d.get('m1'), d.get('m2'), d.get('theta'),
                                 d.get('vx'), d.get('vy'),
                                 d.get('mx_bot'), d.get('my_bot'),
                                 d.get('mx_top'), d.get('my_top')])
        _tri_rows(results.get('analysis_cases'), "Analysis")
        _tri_rows(results.get('combinations'), "Combination")
        T.append(("Area forces",
                  ["Case type", "Case", "Area",
                   "mx [kNm/m]", "my [kNm/m]", "mxy [kNm/m]",
                   "m1 [kNm/m]", "m2 [kNm/m]", "θ [°]",
                   "vx [kN/m]", "vy [kN/m]",
                   "m*x,bot", "m*y,bot", "m*x,top", "m*y,top"], rows))
    else:
        def _tri_rows(container, kind):
            for case, cr in (container or {}).items():
                if not isinstance(cr, dict):
                    continue
                for tid, d in (cr.get('tri_stress') or {}).items():
                    if sc_areas is not None and tid not in sc_areas:
                        continue
                    rows.append([kind, case, tid, d.get('formulation', ''),
                                 d.get('sx'), d.get('sy'), d.get('txy'),
                                 d.get('s1'), d.get('s2'), d.get('theta'), d.get('vm')])
        _tri_rows(results.get('analysis_cases'), "Analysis")
        _tri_rows(results.get('combinations'), "Combination")
        T.append(("Panel stresses",
                  ["Case type", "Case", "Panel", "Formulation",
                   "σx [kN/m²]", "σy [kN/m²]", "τxy [kN/m²]",
                   "σ1 [kN/m²]", "σ2 [kN/m²]", "θ [°]", "vM [kN/m²]"], rows))

    # ── Reinforcement (EC2 concrete design) ──────────────────────────
    reinf = results.get('reinforcement') or {}
    prefs = reinf.get('prefs') or {}
    if prefs:
        T.append(("Design parameters",
                  ["γc", "αcc", "γs", "Strengths"],
                  [[prefs.get('gamma_c', ''), prefs.get('alpha_cc', ''),
                    prefs.get('gamma_s', ''), "fck / fyk from each material"]]))
    bar_rows = reinf.get('rows') or []
    is_grillage = bool(bar_rows) and bar_rows[0].get('kind') == 'grillage'
    rows = []
    if is_grillage:
        # Grillage bar: out-of-plane bending + shear + St-Venant torsion. The
        # axial column is replaced by the torsion T_Ed, and the extra torsion
        # steel (Asl,tor) and the shear-torsion interaction ratio are shown.
        # When the section opts into perimeter distribution of Asl,tor
        # (Section.rc_torsion_distribution == "perimeter",
        # dev/GRILLAGE_DESIGN.md §5.1/§5.3/Phase 3) some side-face steel is
        # non-zero, so an extra column shows the left/right split -- absent
        # (and the column omitted) for the default 50/50 top/bottom mode.
        has_side_faces = any(
            (d.get('Asl_tor_by_face') or {}).get('side_left') for d in bar_rows)
        for d in bar_rows:
            if sc_bars is not None and d.get('element') not in sc_bars:
                continue
            cot = d.get('cot')
            inter = d.get('interaction')
            row = [d.get('element'), d.get('combination'),
                   d.get('location'),
                   d.get('M_Ed'), d.get('V_Ed'), d.get('T_Ed'),
                   (d.get('As_bot') or 0.0) * 1e4,
                   (d.get('As_top') or 0.0) * 1e4,
                   (d.get('Asw_s') or 0.0) * 1e4,
                   (d.get('Asl_tor') or 0.0) * 1e4]
            if has_side_faces:
                faces = d.get('Asl_tor_by_face') or {}
                sl = (faces.get('side_left') or 0.0) * 1e4
                sr = (faces.get('side_right') or 0.0) * 1e4
                row.append(f"{sl:.2f} / {sr:.2f}")
            row += [(f"{cot:.1f}" if isinstance(cot, (int, float)) else ""),
                    (f"{inter:.2f}" if isinstance(inter, (int, float))
                     else ""),
                    "CRUSHING" if d.get('crushing') else "",
                    "★" if d.get('governing') else ""]
            rows.append(row)
        headers = ["Element", "Combination", "Location", "M_Ed [kNm]",
                   "V_Ed [kN]", "T_Ed [kNm]", "As,bot [cm²]", "As,top [cm²]",
                   "Asw/s [cm²/m]", "Asl,tor [cm²]"]
        if has_side_faces:
            headers.append("Asl,tor sides L/R [cm²]")
        headers += ["cot θ", "T/TRd+V/VRd", "Strut", "Gov."]
        T.append(("Grillage reinforcement (EC2 bending + shear + torsion)",
                  headers, rows))
    else:
        for d in bar_rows:
            if sc_bars is not None and d.get('element') not in sc_bars:
                continue
            cot = d.get('cot')
            rows.append([d.get('element'), d.get('combination'),
                         d.get('location'), d.get('reason', ''),
                         d.get('M_Ed'), d.get('N_Ed'), d.get('V_Ed'),
                         (d.get('As_bot') or 0.0) * 1e4,
                         (d.get('As_top') or 0.0) * 1e4,
                         (d.get('Asw_s') or 0.0) * 1e4,
                         (f"{cot:.1f}" if isinstance(cot, (int, float)) else ""),
                         "CRUSHING" if d.get('crushing') else "",
                         "★" if d.get('governing') else "",
                         d.get('note') or ""])
        headers = ["Element", "Combination", "Location", "Reason", "M_Ed [kNm]",
                   "N_Ed [kN]", "V_Ed [kN]", "As,bot [cm²]", "As,top [cm²]",
                   "Asw/s [cm²/m]", "cot θ", "Strut", "Gov.", "Note"]
        if not any(r[-1] for r in rows):     # no flexure notes: drop the column
            headers.pop()
            rows = [r[:-1] for r in rows]
        T.append(("Reinforcement (EC2 design)", headers, rows))

    # ── Triangle reinforcement (EC2 design) ──────────────────────────
    # A plate model designs its slabs for bending (Wood-Armer + EC2 flexure);
    # a plane model designs its membranes (Wood/Baumann). The two report
    # different demands (moments vs membrane forces), so the table adapts to
    # whichever the rows carry.
    tri_reinf = results.get('tri_reinforcement') or {}
    tri_rows = tri_reinf.get('rows') or []
    is_slab = bool(tri_rows) and tri_rows[0].get('kind') == 'slab'
    rows = []
    if is_slab:
        for d in tri_rows:
            if sc_areas is not None and d.get('triangle') not in sc_areas:
                continue
            rows.append([d.get('triangle'), d.get('combination'),
                         d.get('formulation', ''),
                         d.get('mx_bot'), d.get('my_bot'),
                         d.get('mx_top'), d.get('my_top'),
                         (d.get('Asx_bot') or 0.0) * 1e4,
                         (d.get('Asy_bot') or 0.0) * 1e4,
                         (d.get('Asx_top') or 0.0) * 1e4,
                         (d.get('Asy_top') or 0.0) * 1e4,
                         "★" if d.get('governing') else ""])
        T.append(("Slab reinforcement (EC2 flexure)",
                  ["Triangle", "Combination", "Formulation",
                   "m*x,bot [kNm/m]", "m*y,bot [kNm/m]",
                   "m*x,top [kNm/m]", "m*y,top [kNm/m]",
                   "Asx,bot [cm²/m]", "Asy,bot [cm²/m]",
                   "Asx,top [cm²/m]", "Asy,top [cm²/m]", "Gov."],
                  rows))
    else:
        for d in tri_rows:
            if sc_areas is not None and d.get('triangle') not in sc_areas:
                continue
            rows.append([d.get('triangle'), d.get('combination'),
                         d.get('formulation', ''),
                         d.get('n_xx'), d.get('n_yy'), d.get('n_xy'),
                         (d.get('Asx_bot') or 0.0) * 1e4,
                         (d.get('Asx_top') or 0.0) * 1e4,
                         (d.get('Asy_bot') or 0.0) * 1e4,
                         (d.get('Asy_top') or 0.0) * 1e4,
                         # σc = Nc/t: eurocodepy returns the concrete
                         # compression force per unit length; the stress is what
                         # fcd compares against.
                         (d.get('sigma_c') if d.get('sigma_c') is not None
                          else (d.get('Nc') or 0.0) / (d.get('thickness') or 1.0)
                          ) / 1e3,
                         (d.get('fcd') or 0.0) / 1e3,
                         "CRUSHING" if d.get('crushing') else "",
                         "★" if d.get('governing') else ""])
        T.append(("Shell reinforcement (EC2 design)",
                  ["Triangle", "Combination", "Formulation",
                   "n_xx [kN/m]", "n_yy [kN/m]", "n_xy [kN/m]",
                   "Asx,bot [cm²/m]", "Asx,top [cm²/m]",
                   "Asy,bot [cm²/m]", "Asy,top [cm²/m]",
                   "σc [MPa]", "fcd [MPa]", "Strut", "Gov."],
                  rows))

    # ── Punching shear (EC2 §6.4) ─────────────────────────────────────
    punch = results.get('punching') or {}
    rows = []
    for d in (punch.get('rows') or []):
        if sc_nodes is not None and d.get('node') not in sc_nodes:
            continue
        uout = d.get('u_out_eff')
        verdict = ("CRUSHING" if d.get('crushing')
                   else ("REINF NEEDED" if d.get('needs_reinf') else "OK"))
        rows.append([d.get('column'), d.get('node'), d.get('combination'),
                     d.get('position'), d.get('N_Ed'), d.get('beta'),
                     (d.get('d') or 0.0) * 1e3,          # d [m] → mm
                     d.get('u1'), d.get('v_Ed'), d.get('v_Rdc'),
                     d.get('v_Rd_max'), d.get('utilization'),
                     (f"{uout:.3f}" if uout is not None else "—"),
                     verdict,
                     "★" if d.get('governing') else ""])
    if rows:
        T.append(("Punching shear (EN 1992-1-1:2023)",
                  ["Column", "Node", "Combination", "Position", "N_Ed [kN]",
                   "β", "d [mm]", "u1 [m]", "τEd [MPa]", "τRd,c [MPa]",
                   "τRd,max [MPa]", "τEd/τRd,c", "u_out,ef [m]",
                   "Verdict", "Gov."],
                  rows))

    # ── Crack width (EC2 SLS verification) ───────────────────────────
    crack = results.get('crack') or {}
    rows = []
    sc_bars_areas = (None if sc_bars is None and sc_areas is None
                     else (sc_bars or set()) | (sc_areas or set()))
    for d in (crack.get('rows') or []):
        # "Element" here is a bar id (EC2 beam/grid design) or a triangle/quad
        # id (slab/membrane design) depending on the model — test against the
        # union of both scopes rather than assuming one kind.
        if sc_bars_areas is not None and d.get('element') not in sc_bars_areas:
            continue
        ok = d.get('ok')
        if d.get('asmin'):
            check = "ASMIN"
        else:
            check = ("OK" if ok else "NOT OK") if ok is not None else ""
        rows.append([d.get('element'), d.get('combination'), d.get('location'),
                     d.get('M_Ed'), d.get('N_Ed'),
                     "yes" if d.get('cracked', True) else "no",
                     d.get('wk'),
                     d.get('wk_limit'),
                     check,
                     "★" if d.get('governing') else ""])
    T.append(("Crack width (SLS)",
              ["Element", "Combination", "Location", "M_Ed [kNm]", "N_Ed [kN]",
               "Cracked", "wk [mm]", "w_max [mm]", "Check", "Gov."],
              rows))
    return [t for t in T if t[2]]


def _tables_to_lines(tables, max_width: int = 116):
    """Render [(title, header, rows)] as aligned monospaced text lines.

    Numbers use fixed decimals per column (from the header unit) and are
    right-aligned so decimal points line up; text columns stay left-aligned.
    A table wider than *max_width* characters is split into column groups that
    repeat the first (key) column, each labelled "(cont.)", so nothing is lost
    off the right edge.
    """
    L = []
    for title, header, rows in tables:
        cols = [str(h) for h in header]
        decs = [_header_decimals(h) for h in cols]
        numeric = _numeric_columns(cols, rows)
        data = [[_fmt_cell(v, decs[i]) for i, v in enumerate(r)] for r in rows]
        widths = [max(len(cols[i]), max((len(d[i]) for d in data), default=0))
                  for i in range(len(cols))]

        # Split columns into groups that fit max_width, repeating column 0.
        groups = _split_column_groups(widths, max_width, key=0)
        for gi, idxs in enumerate(groups):
            suffix = "" if gi == 0 else "  (cont.)"
            L += ['=' * 50, f"  {title.upper()}{suffix}", '=' * 50]

            def fmt(vals):
                parts = []
                for i in idxs:
                    cell = vals[i]
                    parts.append(cell.rjust(widths[i]) if i in numeric
                                 else cell.ljust(widths[i]))
                return "  " + "  ".join(parts)

            L.append(fmt(cols))
            L.append("  " + "-" * (sum(widths[i] for i in idxs)
                                   + 2 * (len(idxs) - 1)))
            for d in data:
                L.append(fmt(d))
            L.append("")
    return L


def _split_column_groups(widths, max_width, key=0):
    """Group column indices so each group's total printed width ≤ max_width,
    repeating the *key* column at the head of every group after the first.
    Returns a list of index lists (always covering every column once)."""
    n = len(widths)
    sep = 2

    def gwidth(idxs):
        return sum(widths[i] for i in idxs) + sep * (len(idxs) + 1)

    # If the whole row fits, keep it in one group.
    if gwidth(list(range(n))) <= max_width:
        return [list(range(n))]
    groups, cur = [], ([key] if key is not None else [])
    for i in range(n):
        if i == key:
            continue
        trial = cur + [i]
        if cur and gwidth(trial) > max_width and cur != [key]:
            groups.append(cur)
            cur = ([key] if key is not None else []) + [i]
        else:
            cur = trial
    if cur and cur != [key]:
        groups.append(cur)
    return groups or [list(range(n))]


# ── Report section ordering (main body vs annex) ────────────────────────────
# Short, per-model summary tables belong in the main body; long, row-per-entity
# enumerations (nodes, elements, per-node/element results) go to the annex.
_MAIN_TITLES = {"Project", "Model info", "Materials",
                "Load cases", "Analysis cases", "Combinations", "Reaction sums"}
# Order the main-body tables read in.
_MAIN_ORDER = ["Project", "Model info", "Materials", "Sections",
               "Panel sections", "Slab sections", "CST sections",
               "Load cases", "Analysis cases", "Combinations", "Reaction sums"]

# Curated column subsets for the widest summary tables (the full set stays in
# the .xlsx). Keyed by title → list of header labels to keep, in order.
_CURATED_COLUMNS = {
    "Materials": ["Name", "Type", "E [kN/m²]", "Poisson ν", "γ [kN/m³]",
                  "α [1/°C]"],
    "Sections": ["Name", "Type", "Shape", "Material", "b [m]", "h [m]",
                 "Area [m²]", "I [m⁴]"],
}


def _is_annex(title: str) -> bool:
    """True for the long, row-heavy tables that go to the annex."""
    if title in _MAIN_TITLES:
        return False
    if title.endswith("sections"):        # Panel / Slab / CST sections
        return False
    return True


def _curate(title, header, rows):
    """Drop non-essential columns of a wide summary table for the report; the
    full table is still written to the .xlsx. Unknown titles pass through."""
    keep = _CURATED_COLUMNS.get(title)
    if not keep:
        return header, rows
    idx = [header.index(h) for h in keep if h in header]
    return ([header[i] for i in idx],
            [[r[i] for i in idx] for r in rows])


def _toc_lines(entries) -> list:
    """A 'Contents' block (monospaced, no page numbers) for the PDF report."""
    L = ['=' * 50, "  CONTENTS", '=' * 50]
    for i, e in enumerate(entries, 1):
        L.append(f"  {i:>2}.  {e}")
    L.append("")
    return L


def _order_main_tables(tables):
    """Sort main-body tables into the report reading order (_MAIN_ORDER),
    keeping any unlisted ones after, in their given order."""
    pos = {t: i for i, t in enumerate(_MAIN_ORDER)}
    return sorted(tables, key=lambda t: pos.get(t[0], len(pos)))


def split_report_tables(model_tables, result_tables):
    """Partition selected tables into (main, annex), each curated and ordered.

    Main body = summary tables (curated); annex = row-heavy enumerations. Result
    summary rows (Reaction sums) join the main body; the rest of the result
    tables go to the annex.
    """
    main = []
    for (t, h, r) in model_tables:
        if not _is_annex(t):
            ch, cr = _curate(t, h, r)
            main.append((t, ch, cr))
    for (t, h, r) in result_tables:
        if not _is_annex(t):
            main.append((t, h, r))
    main = _order_main_tables(main)
    annex = ([tbl for tbl in model_tables if _is_annex(tbl[0])]
             + [tbl for tbl in result_tables if _is_annex(tbl[0])])
    return main, annex


def _xlsx_from_tables(path, tables):
    """Write [(title, header, rows)] to an .xlsx, one sheet per table."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    bold = Font(bold=True)
    seen = {}
    for title, header, rows in tables:
        name = title[:31]
        if name in seen:
            seen[name] += 1; name = f"{title[:28]}_{seen[name]}"
        else:
            seen[name] = 1
        ws = wb.create_sheet(name)
        ws.append(list(header))
        for c in ws[1]:
            c.font = bold
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[letter].width = min(max(w + 2, 9), 60)
    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])
    if not wb.sheetnames:
        wb.create_sheet("Export")
    wb.save(path)


def _docx_from(path, doc_title, blocks):
    """Build a .docx. blocks is an ordered list of:
       ('heading', text) | ('table', title, header, rows) | ('image', caption, png)."""
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading(doc_title, level=0)
    for blk in blocks:
        kind = blk[0]
        if kind == 'toc':
            doc.add_heading("Contents", level=1)
            for i, e in enumerate(blk[1], 1):
                doc.add_paragraph(f"{i}.  {e}")
        elif kind == 'heading':
            doc.add_heading(blk[1], level=1)
        elif kind == 'image':
            _, caption, png = blk
            if caption:
                doc.add_heading(caption, level=1)
            doc.add_picture(png, width=Inches(6.3))
        elif kind == 'table':
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            _, title, header, rows = blk
            doc.add_heading(title, level=1)
            cols = [str(h) for h in header]
            decs = [_header_decimals(h) for h in cols]
            numeric = _numeric_columns(cols, rows)
            tbl = doc.add_table(rows=1, cols=len(cols))
            tbl.style = 'Light Grid Accent 1'
            tbl.autofit = True
            tbl.allow_autofit = True
            for j, h in enumerate(cols):
                cell = tbl.rows[0].cells[j]
                cell.text = h
                for p in cell.paragraphs:
                    if j in numeric:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for run in p.runs:
                        run.bold = True
            for r in rows:
                cells = tbl.add_row().cells
                for j, val in enumerate(r):
                    cells[j].text = _fmt_cell(val, decs[j])
                    if j in numeric:
                        for p in cells[j].paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(path)


def _iter_result_cases(results: dict):
    """Yield (kind, case, displacements, reactions, element_forces) for every
    combination (envelopes split into MAX/MIN) and analysis case.

    Load cases are omitted — they are load definitions only; their solved
    response is reported via the matching Linear analysis cases."""
    for cid, cres in results.get('combinations', {}).items():
        d = cres.get('displacements', {}); r = cres.get('reactions', {})
        e = cres.get('element_forces', {})
        if isinstance(d, dict) and 'max' in d:
            yield ("Combination", f"{cid} (MAX)", d.get('max', {}),
                   r.get('max', {}), e.get('max', {}))
            yield ("Combination", f"{cid} (MIN)", d.get('min', {}),
                   r.get('min', {}), e.get('min', {}))
        else:
            yield ("Combination", cid, d, r, e)
    for ac, acr in results.get('analysis_cases', {}).items():
        if not isinstance(acr, dict) or 'error' in acr or 'modal_info' in acr:
            continue
        if 'displacements' in acr:
            yield ("Analysis case", ac, acr.get('displacements', {}),
                   acr.get('reactions', {}), acr.get('element_forces', {}))


# ---------------------------------------------------------------------------
# Plain-text results report (shared by the GUI results panel and exports)
# ---------------------------------------------------------------------------


def _fmt_ef(label: str, ef: dict, lines: list):
    """Append formatted element end-force lines."""
    for eid, v in ef.items():
        if isinstance(v, dict) and 'i' in v:
            ni, nj = v['i'], v['j']
            lines.append(
                f"    {eid}  i: N={ni[0]:+.4f}  V={ni[1]:+.4f}  M={ni[2]:+.4f}")
            lines.append(
                f"    {eid}  j: N={nj[0]:+.4f}  V={nj[1]:+.4f}  M={nj[2]:+.4f}")


def _fmt_disp(disp: dict, lines: list):
    for nid, d in disp.items():
        if isinstance(d, list):
            lines.append(
                f"    {nid}: ux={d[0]:+.6f}  uy={d[1]:+.6f}  rz={d[2]:+.6f}")


def _fmt_reac(reac: dict, lines: list, show_sum: bool = True):
    sx = sy = sm = 0.0
    n = 0
    for nid, r in reac.items():
        if isinstance(r, list):
            lines.append(
                f"    {nid}: Rx={r[0]:+.4f}  Ry={r[1]:+.4f}  Mz={r[2]:+.4f}")
            sx += r[0]; sy += r[1]; sm += r[2]; n += 1
    # Sum of all reactions = total support force (an equilibrium check: it should
    # balance the applied loads). Skipped for envelopes, where per-node extremes
    # don't sum to a physical total.
    if show_sum and n:
        lines.append(
            f"    Σ (all supports): Rx={sx:+.4f}  Ry={sy:+.4f}  Mz={sm:+.4f}")


def _build_results_lines(results: dict) -> list:
    """Build the full plain-text results report (shared by the panel + exports)."""
    lines = []

    # Load-case results are intentionally omitted — load cases are load
    # definitions only; their solved response is reported via the matching
    # Linear analysis cases instead.

    # ── Combinations ─────────────────────────────────────────
    for combo_id, combo_res in results.get('combinations', {}).items():
        lines.append(f"{'='*50}")
        lines.append(f"  COMBINATION: {combo_id}")
        lines.append(f"{'='*50}")

        disp_res = combo_res.get('displacements', {})
        reac_res = combo_res.get('reactions', {})
        ef_res   = combo_res.get('element_forces', {})

        # Envelope stores max/min sub-dicts
        if isinstance(disp_res, dict) and 'max' in disp_res:
            lines.append("  Displacements — MAX (m, rad):")
            _fmt_disp(disp_res['max'], lines)
            lines.append("  Displacements — MIN (m, rad):")
            _fmt_disp(disp_res['min'], lines)
            lines.append("  Reactions — MAX (kN, kNm):")
            _fmt_reac(reac_res.get('max', {}), lines, show_sum=False)
            lines.append("  Reactions — MIN (kN, kNm):")
            _fmt_reac(reac_res.get('min', {}), lines, show_sum=False)
            lines.append("  Element forces — MAX (kN, kNm):")
            _fmt_ef(combo_id, ef_res.get('max', {}), lines)
            lines.append("  Element forces — MIN (kN, kNm):")
            _fmt_ef(combo_id, ef_res.get('min', {}), lines)
        else:
            lines.append("  Displacements (m, rad):")
            _fmt_disp(disp_res, lines)
            lines.append("  Reactions (kN, kNm):")
            _fmt_reac(reac_res, lines)
            lines.append("  Element end-forces (kN, kNm):")
            _fmt_ef(combo_id, ef_res, lines)

            # Span distribution extremes
            dist = results.get('combo_distribution', {}).get(combo_id, {})
            if dist:
                lines.append("  Element span extremes  [N / V / M]  (kN, kNm):")
                for eid, d in dist.items():
                    Nmax = float(np.max(np.abs(d['N'])))
                    Vmax = float(np.max(np.abs(d['V'])))
                    Mpos = float(np.max(d['M']))
                    Mneg = float(np.min(d['M']))
                    lines.append(
                        f"    {eid}: |N|max={Nmax:.4f}  |V|max={Vmax:.4f}"
                        f"  M+={Mpos:.4f}  M-={Mneg:.4f}")
        lines.append("")

    # ── Analysis cases ────────────────────────────────────────
    for ac_id, ac_res in results.get('analysis_cases', {}).items():
        if not ac_res:
            continue
        if 'error' in ac_res:
            lines.append(f"{'='*50}")
            lines.append(f"  ANALYSIS CASE: {ac_id}  [ERROR]")
            lines.append(f"  {ac_res['error']}")
            lines.append("")
            continue

        # Modal results
        if 'modal_info' in ac_res:
            lines.append(f"{'='*50}")
            lines.append(f"  MODAL ANALYSIS: {ac_id}")
            lines.append(f"{'='*50}")
            lines.append(f"  Total mass X: {ac_res.get('total_mass_x', 0):.4f} t"
                         f"   Y: {ac_res.get('total_mass_y', 0):.4f} t")
            lines.append("")
            lines.append(f"  {'Mode':>4}  {'ω [rad/s]':>12}  {'f [Hz]':>10}"
                         f"  {'T [s]':>10}  {'Meff_X%':>8}  {'Meff_Y%':>8}"
                         f"  {'ΣX%':>7}  {'ΣY%':>7}")
            lines.append("  " + "-"*80)
            cum_x = cum_y = 0.0
            for mi in ac_res['modal_info']:
                cum_x += mi['meff_x_pct']
                cum_y += mi['meff_y_pct']
                lines.append(
                    f"  {mi['mode']:>4}  {mi['omega']:>12.4f}  {mi['frequency']:>10.4f}"
                    f"  {mi['period']:>10.4f}  {mi['meff_x_pct']:>7.1f}%"
                    f"  {mi['meff_y_pct']:>7.1f}%  {cum_x:>6.1f}%  {cum_y:>6.1f}%")
            lines.append(f"  {'Cumulative':>36}  {cum_x:>7.1f}%  {cum_y:>7.1f}%")
            lines.append("")

        # Mass results
        elif 'nodal_masses' in ac_res:
            lines.append(f"{'='*50}")
            lines.append(f"  MASS CASE: {ac_id}")
            lines.append(f"{'='*50}")
            lines.append(f"  Total mass: {ac_res.get('total_mass', 0):.4f} t")
            for nid, m in ac_res['nodal_masses'].items():
                if isinstance(m, dict):
                    lines.append(f"    {nid}: mx={m['mx']:.4f}  my={m['my']:.4f}"
                                 f"  mtz={m['mtz']:.4f} t·m²")
                else:
                    lines.append(f"    {nid}: {m:.4f} t")
            lines.append("")

        # Spectrum analysis case (peak absolute values)
        elif ac_res.get('is_spectrum') and 'displacements' in ac_res:
            rule = ac_res.get('combination_rule', '')
            dirn = ac_res.get('direction', '')
            lines.append(f"{'='*50}")
            lines.append(f"  SPECTRUM: {ac_id}  [{rule} / {dirn}]  (peak absolute values)")
            lines.append(f"{'='*50}")
            lines.append("  Peak displacements (m, rad):")
            _fmt_disp(ac_res['displacements'], lines)
            lines.append("  Peak reactions (kN, kNm):")
            _fmt_reac(ac_res['reactions'], lines)
            lines.append("  Peak element end-forces (kN, kNm):")
            _fmt_ef(ac_id, ac_res['element_forces'], lines)
            lines.append("")

        # Linear analysis case
        elif 'displacements' in ac_res:
            lines.append(f"{'='*50}")
            lines.append(f"  ANALYSIS CASE: {ac_id}")
            lines.append(f"{'='*50}")
            lines.append("  Displacements (m, rad):")
            _fmt_disp(ac_res['displacements'], lines)
            lines.append("  Reactions (kN, kNm):")
            _fmt_reac(ac_res['reactions'], lines)
            lines.append("  Element end-forces (kN, kNm):")
            _fmt_ef(ac_id, ac_res['element_forces'], lines)
            dist = ac_res.get('element_distribution', {})
            if dist:
                lines.append("  Element span extremes  [N / V / M]  (kN, kNm):")
                for eid, d in dist.items():
                    Nmax = float(np.max(np.abs(d['N'])))
                    Vmax = float(np.max(np.abs(d['V'])))
                    Mpos = float(np.max(d['M']))
                    Mneg = float(np.min(d['M']))
                    lines.append(
                        f"    {eid}: |N|max={Nmax:.4f}  |V|max={Vmax:.4f}"
                        f"  M+={Mpos:.4f}  M-={Mneg:.4f}")
            lines.append("")

    # ── Reinforcement (EC2 concrete design) ──────────────────
    reinf = results.get('reinforcement') or {}
    rrows = reinf.get('rows') or []
    if rrows:
        prefs = reinf.get('prefs') or {}
        lines.append(f"{'='*50}")
        lines.append("  REINFORCEMENT (EC2 concrete design)")
        lines.append(f"{'='*50}")
        if prefs:
            lines.append(
                f"  γc={prefs.get('gamma_c', '')}, αcc={prefs.get('alpha_cc', '')}, "
                f"γs={prefs.get('gamma_s', '')}  (fck/fyk from each section's material)")
        lines.append(f"  {'Elem':>6}  {'Combination':>16}  {'Loc':>8}"
                     f"  {'As,bot':>8}  {'As,top':>8}  {'Asw/s':>8}  Gov")
        lines.append("  " + "-"*70)
        for d in rrows:
            lines.append(
                f"  {str(d.get('element','')):>6}  {str(d.get('combination','')):>16}"
                f"  {str(d.get('location','')):>8}"
                f"  {(d.get('As_bot') or 0)*1e4:>7.2f}  {(d.get('As_top') or 0)*1e4:>7.2f}"
                f"  {(d.get('Asw_s') or 0)*1e4:>7.2f}"
                f"  {'★' if d.get('governing') else ''}")
        lines.append("  (As in cm², Asw/s in cm²/m)")
        lines.append("")

    return lines
