# Copyright (c) 2026 Paulo Cachim
# SPDX-License-Identifier: LGPL-3.0-or-later

"""Render design calculation reports (eurocodepy ``CalcReport`` dicts).

The design routines (``design_steel_members``, ``design_timber_members``,
``design_concrete_sections``, ``design_punching``) can, with ``with_reports``,
attach a full step-by-step report — a :class:`eurocodepy.calc_report.CalcReport`
``to_dict()`` — for the governing case of each member / section / column. This
module turns those report dicts into human output:

* :func:`render_markdown` — plain Markdown (debug / preview);
* :func:`render_html` — a self-contained HTML page with MathJax for the LaTeX;
* :func:`render_docx` — a Word document (needs ``python-docx``).

It is **UI-agnostic and Qt-free**: the GUI collects the reports (running a design
with ``with_reports=True``) and calls one of these renderers. A report dict has
the shape::

    {schema_version, title, meta:{...},
     sections:[{title, steps:[{symbol,value,unit,clause,expr,latex,subst,
                               note,ok}]}]}
"""

from __future__ import annotations

import html as _html

# The report-dict schema this renderer understands. Keep in sync with
# eurocodepy.calc_report.SCHEMA_VERSION; a mismatch is surfaced, not fatal.
EXPECTED_SCHEMA = 1


def schema_mismatches(reports: dict) -> list:
    """Ids whose report was produced under a different schema version."""
    return [rid for rid, rep in reports.items()
            if int(rep.get("schema_version", EXPECTED_SCHEMA)) != EXPECTED_SCHEMA]


# ── font choices (Preferences ▸ Design ▸ Design report) ─────────────────────
#
# A curated list, not a free system-font picker: the PDF's equations render
# through matplotlib's mathtext, which only understands a handful of built-in
# "font sets" (dejavusans/dejavuserif/cm/stix/stixsans) — it cannot follow an
# arbitrary installed font the way plain text can. Each entry here pairs a
# mathtext font set with a matching plain-text family/CSS/Word font, so
# headings, body text and equations read as one consistent typeface in every
# export format instead of mixing an arbitrary body font with mathtext's
# default. ``DEFAULT_REPORT_FONT_FAMILY``/``DEFAULT_REPORT_FONT_SIZE`` match
# what every renderer used before this became a preference.
REPORT_FONT_CHOICES = {
    "sans": {
        "label": "Sans (Arial / DejaVu Sans)",
        "mpl_family": "sans-serif", "mathtext_fontset": "dejavusans",
        "css_family": "Arial, Helvetica, sans-serif", "docx_name": "Arial",
    },
    "serif": {
        "label": "Serif (Georgia / DejaVu Serif)",
        "mpl_family": "serif", "mathtext_fontset": "dejavuserif",
        "css_family": "Georgia, 'Times New Roman', serif", "docx_name": "Georgia",
    },
    "stix": {
        "label": "Times-like (STIX)",
        "mpl_family": "serif", "mathtext_fontset": "stix",
        "css_family": "'Times New Roman', Times, serif",
        "docx_name": "Times New Roman",
    },
    "cm": {
        "label": "Computer Modern (LaTeX look)",
        "mpl_family": "serif", "mathtext_fontset": "cm",
        "css_family": "'Latin Modern Roman', 'CMU Serif', serif",
        "docx_name": "Cambria Math",
    },
}
DEFAULT_REPORT_FONT_FAMILY = "sans"
DEFAULT_REPORT_FONT_SIZE = 10


def _report_font_choice(font_family: str | None) -> dict:
    return REPORT_FONT_CHOICES.get(font_family or DEFAULT_REPORT_FONT_FAMILY,
                                   REPORT_FONT_CHOICES[DEFAULT_REPORT_FONT_FAMILY])


# ── normalisation ───────────────────────────────────────────────────────────

def reports_from(design_output) -> dict:
    """Extract ``{id: report_dict}`` from a design routine's output.

    Accepts either a dict with a ``"reports"`` key (steel / timber) or a list of
    rows carrying a ``"report"`` on the governing ones (concrete / punching).
    """
    if isinstance(design_output, dict):
        return dict(design_output.get("reports", {}))
    out: dict = {}
    for row in design_output or []:
        rep = row.get("report") if isinstance(row, dict) else None
        if rep:
            key = (row.get("element") or row.get("column") or row.get("member")
                   or row.get("triangle") or str(len(out)))
            out[key] = rep
    return out


def build_reports(struc, results, material: str, *, prefs: dict | None = None,
                  combinations=None) -> dict:
    """Run the design for *material* with reports on and return ``{id: report}``.

    ``material`` is one of ``"steel"``, ``"timber"``, ``"concrete"``,
    ``"punching"``. ``prefs`` is the design-preferences dict (the same one the
    normal design uses); the concrete / punching partial factors are read from
    it. This is the single Qt-free seam the GUI calls before rendering.
    """
    prefs = prefs or {}
    gc = float(prefs.get("gamma_c", 1.5))
    gs = float(prefs.get("gamma_s", 1.15))
    acc = float(prefs.get("alpha_cc", 1.0))
    if material == "steel":
        from .steel_design import design_steel_members
        return reports_from(design_steel_members(
            struc, results, prefs=prefs, combinations=combinations,
            with_reports=True))
    if material == "timber":
        from .timber_design import design_timber_members
        return reports_from(design_timber_members(
            struc, results, prefs=prefs, combinations=combinations,
            with_reports=True))
    if material == "concrete":
        from .rc_design import design_concrete_sections
        return reports_from(design_concrete_sections(
            struc, results, gamma_c=gc, gamma_s=gs, alpha_cc=acc,
            combinations=combinations, with_reports=True))
    if material == "concrete_area":
        # Concrete *areas*: slabs (plate domain, Wood-Armer bending) or
        # membranes (plane domain, Wood/Baumann). Dispatch by domain, mirroring
        # design_tri_and_store.
        if getattr(struc, "domain", "plane") == "plate":
            from .rc_design import design_concrete_slabs
            rows = design_concrete_slabs(
                struc, results, gamma_c=gc, gamma_s=gs, alpha_cc=acc,
                combinations=combinations, with_reports=True)
        else:
            from .rc_design import design_concrete_planes
            rows = design_concrete_planes(
                struc, results, gamma_c=gc, gamma_s=gs, alpha_cc=acc,
                combinations=combinations, with_reports=True)
        return reports_from(rows)
    if material == "punching":
        from .punching import design_punching
        return reports_from(design_punching(
            struc, results, combinations=combinations,
            gamma_c=gc, gamma_s=gs, alpha_cc=acc,
            gamma_v=float(prefs.get("gamma_v", 1.4)),
            eta_sys=float(prefs.get("eta_sys", 1.5)),
            edition=str(prefs.get("ec2_edition", "2004")),
            with_reports=True))
    raise ValueError(f"unknown material {material!r}")


def filter_reports(reports: dict, material: str, struc,
                   selected_bars=None, selected_nodes=None,
                   selected_tris=None) -> dict:
    """Keep only the reports touching the current selection.

    Report keys are member ids (steel/timber), element/bar ids (concrete),
    column ids (punching) or triangle ids (concrete areas). Steel/timber members
    are kept when any of their bars is selected; concrete by bar id; punching by
    the column's supported node; concrete areas by triangle id. An empty/None
    selection returns *reports* unchanged.
    """
    bars = set(selected_bars or [])
    nodes = set(selected_nodes or [])
    tris = set(selected_tris or [])
    if not bars and not nodes and not tris:
        return reports
    if material == "steel":
        # Steel reports mix two kinds under one dict (see
        # design_steel_members): §6.3.3 member-buckling reports keyed by
        # member id, and §6.2 cross-section reports keyed by bar id
        # directly (meta has "element"). Disambiguate per report via its
        # own meta rather than assuming every key is a member id.
        from .member_utils import identify_members
        _members, by_bar = identify_members(struc)
        keep_members = {by_bar[b].id for b in bars if b in by_bar}
        out = {}
        for k, v in reports.items():
            meta = v.get("meta", {}) if isinstance(v, dict) else {}
            if "element" in meta:
                if meta["element"] in bars:
                    out[k] = v
            elif k in keep_members:
                out[k] = v
        return out
    if material in ("timber", "column"):
        from .member_utils import identify_members
        _members, by_bar = identify_members(struc)
        keep = {by_bar[b].id for b in bars if b in by_bar}
        return {k: v for k, v in reports.items() if k in keep}
    if material == "concrete":
        return {k: v for k, v in reports.items() if k in bars}
    if material == "concrete_area":
        return {k: v for k, v in reports.items() if k in tris} if tris else reports
    if material == "punching":
        cols = {c.id for c in getattr(struc, "punch_columns", [])
                if getattr(c, "node_id", None) in nodes}
        return {k: v for k, v in reports.items() if k in cols} if nodes else reports
    return reports


# ── per-combination detail (Qt-free helpers used by the report panel) ────────

def report_base_id(report: dict) -> str:
    """The base member / element / column id a report is for.

    Read from the report ``meta``. Every report writer now sets a material-
    agnostic ``id`` key; the material-specific key (``member`` for
    steel/timber, ``element``/``triangle`` for concrete, ``column`` for
    punching) is kept alongside it for readability and is used as a fallback
    for any report built before ``id`` was introduced. Lets the UI recover the
    canvas target of a report even when the report is keyed by a composite
    per-combination key.
    """
    meta = report.get("meta", {}) if isinstance(report, dict) else {}
    return str(meta.get("id") or meta.get("member") or meta.get("element")
               or meta.get("column") or meta.get("triangle") or "")


def summary_rows(reports: dict, material: str | None = None) -> list:
    """One row per report, read straight from each report's ``meta``.

    Most materials have a real ratio concept (steel/timber utilization
    checks, punching), so their rows are ``{id, section, combination,
    utilization, ok, kind: "ratio"}`` as before. Concrete reinforcement
    sizing has no utilization ratio -- the design routines always set
    ``utilization: None`` -- so a summary built from those fields alone is
    the same tie for every row. For ``material == "concrete"`` (bar
    sections) and ``"concrete_area"`` (tri/quad membranes and slabs) this
    instead surfaces the actual sizing result -- the reinforcement areas
    the design computed (``As_bot``/``As_top``/``Asw_s`` for bars,
    ``Asx_bot``/``Asx_top``/``Asy_bot``/``Asy_top`` for areas), which
    ``rc_design.py`` copies into ``meta`` alongside ``utilization``/``ok``
    -- and sorts worst-first by total steel area, so the governing member
    the row is next to what needs it most rather than an arbitrary order.

    This is the ONLY thing the General tab's summary table is built from:
    since it comes from the exact same ``build_reports()`` call that builds
    the narrative sections, the summary can never show a different number
    than the report it summarises.
    """
    rows = []
    if material == "concrete":
        for rid, rep in reports.items():
            meta = (rep.get("meta") or {}) if isinstance(rep, dict) else {}
            as_bot = meta.get("As_bot"); as_top = meta.get("As_top")
            asw_s = meta.get("Asw_s")
            total = (as_bot or 0.0) + (as_top or 0.0) + (asw_s or 0.0)
            rows.append({
                "id": report_base_id(rep) or str(rid),
                "section": meta.get("profile") or meta.get("section") or "",
                "combination": meta.get("combination", ""),
                "As_bot": as_bot, "As_top": as_top, "Asw_s": asw_s,
                "_total": total, "kind": "concrete_bar",
            })
        rows.sort(key=lambda r: -r["_total"])
        return rows
    if material == "concrete_area":
        for rid, rep in reports.items():
            meta = (rep.get("meta") or {}) if isinstance(rep, dict) else {}
            axb = meta.get("Asx_bot"); axt = meta.get("Asx_top")
            ayb = meta.get("Asy_bot"); ayt = meta.get("Asy_top")
            total = (axb or 0.0) + (axt or 0.0) + (ayb or 0.0) + (ayt or 0.0)
            rows.append({
                "id": report_base_id(rep) or str(rid),
                "section": meta.get("profile") or meta.get("section") or "",
                "combination": meta.get("combination", ""),
                "Asx_bot": axb, "Asx_top": axt, "Asy_bot": ayb, "Asy_top": ayt,
                "_total": total, "kind": "concrete_area",
            })
        rows.sort(key=lambda r: -r["_total"])
        return rows
    for rid, rep in reports.items():
        meta = (rep.get("meta") or {}) if isinstance(rep, dict) else {}
        rows.append({
            "id": report_base_id(rep) or str(rid),
            "section": meta.get("profile") or meta.get("section") or "",
            "combination": meta.get("combination", ""),
            "utilization": meta.get("utilization"),
            "ok": meta.get("ok"),
            "kind": "ratio",
        })
    rows.sort(key=lambda r: (r["utilization"] is None, -(r["utilization"] or 0.0)))
    return rows


def per_combination_key(base_id, combination) -> str:
    """The dict key for a single-combination report of *base_id*."""
    return f"{base_id} · {combination}"


def merge_per_combination(per_combo_reports: dict) -> dict:
    """Flatten ``{combination: {id: report}}`` into ``{"<id> · <combo>": report}``.

    The panel's *Per combination* detail builds one ``{id: report}`` per chosen
    combination (each is the governing case *within* that single combination) and
    merges them here so every (member, combination) pair is one report entry.
    """
    out: dict = {}
    for combo, reports in per_combo_reports.items():
        for rid, rep in (reports or {}).items():
            out[per_combination_key(rid, combo)] = rep
    return out


def _normalize_mathtext(latex: str) -> str:
    """Fix the handful of LaTeX spellings matplotlib mathtext does not accept
    (``\\le`` instead of ``\\leq``) — shared by every mathtext entry point
    (:func:`latex_png`, the native PDF equation renderer, the extent probe)."""
    return latex.replace(r"\le ", r"\leq ").replace(r"\le1", r"\leq 1")


# A throwaway 1x1 Agg figure, created once and reused by every latex_png()/
# _mathtext_extent_in() call for the lifetime of the process — a design report
# calls these once per equation (dozens to hundreds of times), and going
# through pyplot's plt.figure()/plt.close() cycle for each one measured at
# ~126 ms/call (the global figure-manager bookkeeping, not the actual text
# layout) versus ~0.6 ms/call for a bare Figure reused in place. A report
# with ~100 equations was taking the better part of a minute — synchronously
# on the Qt main thread, with no progress indicator — and looked like the
# app had frozen ("bloqueou"). The reused figure is cleared of its one Text
# artist (``txt.remove()``) right after each use, so nothing accumulates and
# no stale content leaks into the next equation's measurement or PNG.
_scratch_fig = None


def _get_scratch_figure():
    global _scratch_fig
    if _scratch_fig is None:
        import matplotlib
        matplotlib.use("Agg")
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        from matplotlib.figure import Figure

        _scratch_fig = Figure(figsize=(0.01, 0.01))
        FigureCanvasAgg(_scratch_fig)
    return _scratch_fig


def latex_png(latex: str, fontsize: int = 15) -> bytes | None:
    """Typeset a LaTeX math string to a transparent PNG (matplotlib mathtext).

    Returns the PNG bytes, or ``None`` if the expression is empty or mathtext
    cannot parse it (the caller then falls back to plain text). Used for Word
    (no MathJax there) and as the PDF's fallback when the native vector
    renderer (:func:`_mathtext_extent_in` + direct ``Figure.text``, used by
    :func:`render_pdf`) fails to parse an expression; HTML uses MathJax
    directly.
    """
    if not latex:
        return None
    from io import BytesIO
    s = _normalize_mathtext(latex)
    fig = _get_scratch_figure()
    txt = fig.text(0, 0, f"${s}$", fontsize=fontsize)
    try:
        buf = BytesIO()
        fig.savefig(buf, format="png", dpi=150, bbox_inches="tight",
                    pad_inches=0.04, transparent=True)
    except Exception:            # noqa: BLE001 — mathtext parse failure
        return None
    finally:
        txt.remove()
    return buf.getvalue()


def _mathtext_extent_in(s: str, fontsize: int) -> tuple[float, float]:
    """Width/height (inches) that ``$<s>$`` will occupy when drawn with
    :meth:`~matplotlib.figure.Figure.text` at *fontsize* — measured with a
    throwaway figure/renderer, never rendered to a raster. Lets
    :func:`render_pdf` reserve the exact vertical space a **native, vector**
    equation needs, the same way :func:`latex_png` used to size its PNG, but
    without rasterising the formula itself. Raises on a mathtext parse
    failure — the caller falls back to :func:`latex_png`."""
    fig = _get_scratch_figure()
    txt = fig.text(0, 0, f"${s}$", fontsize=fontsize)
    try:
        fig.canvas.draw()
        bbox = txt.get_window_extent()
        dpi = fig.dpi
        return bbox.width / dpi, bbox.height / dpi
    finally:
        txt.remove()


def section_figure_png(geom: dict) -> bytes | None:
    """Draw a cross-section sketch from a ``geom`` dict and return PNG bytes.

    ``geom`` keys: ``shape`` ('rect'|'I'|'RHS'|'SHS'|'CHS'|...), ``b``, ``h``,
    ``tw``, ``tf``, ``d`` (diameter) — all in metres — and an optional
    ``label``. Returns ``None`` when there is not enough geometry to draw.
    """
    if not geom:
        return None
    from io import BytesIO
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Circle, Rectangle

    shape = str(geom.get("shape") or "rect").lower()
    b, h = geom.get("b"), geom.get("h")
    tw, tf, dia = geom.get("tw"), geom.get("tf"), geom.get("d")
    col, ec = "#c8d6f0", "#1f4fd8"
    fig, ax = plt.subplots(figsize=(2.0, 2.0))
    ax.set_aspect("equal"); ax.axis("off")
    try:
        if shape in ("chs", "circular", "circle") and (dia or b):
            r = (dia or b) / 2.0
            ax.add_patch(Circle((0, 0), r, fc=col, ec=ec, lw=2))
            lim = r * 1.3
        elif shape in ("i", "ipe", "he", "hea", "heb") and b and h and tw and tf:
            xs = [-b/2, b/2, b/2, tw/2, tw/2, b/2, b/2, -b/2, -b/2,
                  -tw/2, -tw/2, -b/2, -b/2]
            ys = [h/2, h/2, h/2-tf, h/2-tf, -h/2+tf, -h/2+tf, -h/2, -h/2,
                  -h/2+tf, -h/2+tf, h/2-tf, h/2-tf, h/2]
            ax.fill(xs, ys, fc=col, ec=ec, lw=2)
            lim = max(b, h) * 0.62
        elif shape in ("rhs", "shs", "box") and b and h:
            ax.add_patch(Rectangle((-b/2, -h/2), b, h, fc=col, ec=ec, lw=2))
            t = tw or min(b, h) * 0.1
            ax.add_patch(Rectangle((-b/2+t, -h/2+t), b-2*t, h-2*t,
                                   fc="white", ec=ec, lw=1))
            lim = max(b, h) * 0.62
        elif b and h:
            ax.add_patch(Rectangle((-b/2, -h/2), b, h, fc=col, ec=ec, lw=2))
            ax.text(0, -h/2, f"\nb = {b*1000:.0f} mm", ha="center", va="top",
                    fontsize=8)
            ax.text(-b/2, 0, f"h = {h*1000:.0f} mm  ", ha="right", va="center",
                    rotation=90, fontsize=8)
            lim = max(b, h) * 0.8
        else:
            plt.close(fig)
            return None
        ax.set_xlim(-lim, lim); ax.set_ylim(-lim, lim)
        buf = BytesIO()
        fig.savefig(buf, format="png", dpi=150, bbox_inches="tight",
                    transparent=True)
    except Exception:            # noqa: BLE001
        plt.close(fig)
        return None
    plt.close(fig)
    return buf.getvalue()


def _fmt(v) -> str:
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float):
        return f"{v:.4g}"
    return str(v)


def _fmt_area(v) -> str:
    """A reinforcement area stored in m² (bars) or m²/m (tri/quad areas),
    shown as cm² / cm²/m -- the same 1e4 scaling and unit the canvas
    overlays and per-station callouts already use (see canvas.py's
    ``As,bot``/``As,x`` readouts), so the summary table reads consistently
    with the rest of the app."""
    if v is None:
        return "--"
    return f"{float(v) * 1e4:.2f}"


# Meta keys already shown elsewhere (section drawing / summary table) or too
# internal to mean anything out of context: `reason` is rc_design.py's
# "why this station governs" code (e.g. "Asxb,max"), which for the member-end
# stations is a literal duplicate of `location` ("i"/"j") and elsewhere is
# jargon a report reader has no context for.
_META_EXCLUDE_KEYS = {"section_geom", "utilization", "ok", "reason"}

# Reinforcement areas stored in m² / m²/m -- routed through _fmt_area (cm² /
# cm²/m) instead of the generic _fmt, which would otherwise print them in
# raw scientific notation (e.g. "1.821e-05").
_META_AREA_KEYS = {"As_bot", "As_top", "Asw_s",
                   "Asx_bot", "Asx_top", "Asy_bot", "Asy_top"}


def _meta_display_items(meta: dict) -> list:
    """The ``(key, formatted value)`` pairs to show in a report's meta line --
    shared by the Markdown/HTML/PDF/docx renderers so the same fields are
    dropped/reformatted everywhere instead of drifting across four
    near-duplicate loops."""
    out = []
    for k, v in meta.items():
        if k in _META_EXCLUDE_KEYS:
            continue
        out.append((k, _fmt_area(v) if k in _META_AREA_KEYS else _fmt(v)))
    return out


# Symbols/expressions across eurocodepy's ``trace.step()`` calls use a plain
# "M_Ed", "f_yk", "k_c,y" convention (underscore = subscript) rather than real
# LaTeX ("M_{Ed}"). Outside a math environment a lone "_" is just a literal
# underscore to a Markdown/pandoc renderer, not a subscript — hence the wall
# of stray underscores in the exported report. ``_texify`` turns the
# convention into proper LaTeX grouping and ``_mathify`` wraps it in ``$…$``
# so it renders as an actual subscript wherever the Markdown is viewed.
import re as _re
_SUBSCRIPT_RE = _re.compile(r'([A-Za-zΑ-Ωα-ω][A-Za-z0-9]*)_([A-Za-z0-9,]+)')


def _texify_one(m: "_re.Match") -> str:
    prefix, sub = m.group(1), m.group(2)
    # "My_Ed" / "Vz_Rd" / "As_min": a 2-letter prefix that is itself
    # BASE + one lower-case axis/component letter glued on with no "_" of its
    # own — the axis belongs *inside* the subscript group too:
    # "My_Ed" -> "M_{y,Ed}", not "My_{Ed}". A longer spelled-out name like
    # "chi_y" (prefix "chi", 3 letters) is a single symbol, not base+axis, so
    # it is left as a plain "chi_{y}".
    if len(prefix) == 2 and prefix[0].isupper() and prefix[1].islower():
        return f'{prefix[0]}_{{{prefix[1]},{sub}}}'
    return f'{prefix}_{{{sub}}}'


def _texify(s: str) -> str:
    """"M_Ed" -> "M_{Ed}", "My_Ed" -> "M_{y,Ed}", "f_yk" -> "f_{yk}",
    "k_c,y" -> "k_{c,y}", "chi_y" -> "chi_{y}", ..."""
    return _SUBSCRIPT_RE.sub(_texify_one, s) if s else s


def _mathify(s: str) -> str:
    """``_texify`` then wrap in inline-math ``$…$`` (empty string passes through)."""
    t = _texify(s)
    return f"${t}$" if t else t


def _mathify_bold(s: str) -> str:
    """``_texify`` then wrap as **bold** math via LaTeX's own ``\\mathbf{}``,
    not Markdown's ``**…**``. A ``**`` glued directly onto a ``$`` delimiter
    (``**$M_{Ed}$**``) trips up several Markdown+LaTeX pipelines (pandoc,
    MathJax) — they only reliably render it with a space on both sides
    (``** $M_{Ed}$ **``), which looks wrong in the source. Asking LaTeX for
    the bold instead sidesteps the interaction entirely."""
    t = _texify(s)
    return f"$\\mathbf{{{t}}}$" if t else t


# ── Markdown ────────────────────────────────────────────────────────────────
#
# ``_step_lines(st)`` is the **single source of truth** for how one step is
# worded: both ``render_markdown`` (the .md/.docx-via-pandoc/.html-via-pandoc
# path) and ``render_docx``'s native fallback build the exact same strings
# from it, so the text reads identically regardless of which renderer (or
# whether pandoc is installed) produced the file — only the *formatting*
# differs (Markdown syntax vs. native Word runs via ``_md_line_to_docx``).

def _step_head(st: dict) -> str:
    head = (f"{_mathify_bold(st['symbol'])} = {_fmt(st['value'])} "
           f"{st.get('unit','')}").rstrip()
    if st.get("clause"):
        head += f"  _[{st['clause']}]_"
    return head


def _step_lines(st: dict) -> list[str]:
    """The step's body lines *after* the head line — expr/subst/ok/note, in
    the one order both renderers use. ``latex`` (a display equation) is
    reported separately since it renders differently per format (a ``$$…$$``
    block in Markdown, an image in the pandoc-less docx fallback)."""
    return [line for _kind, line in _step_lines_kv(st)]


def _step_lines_kv(st: dict) -> list[tuple[str, str]]:
    """Same as :func:`_step_lines` but tagged ``(kind, line)`` — ``"expr"``,
    ``"subst"``, ``"ok"`` or ``"not_ok"``, ``"note"`` — so a consumer (PDF
    colouring) can tell them apart without re-parsing the rendered text."""
    lines = []
    if not st.get("latex") and st.get("expr"):
        lines.append(("expr", f"  - {_mathify(st['expr'])}"))
    if st.get("subst"):
        lines.append(("subst", f"  - = {_mathify(st['subst'])}"))
    if st.get("ok") is not None:
        kind = "ok" if st["ok"] else "not_ok"
        lines.append((kind, f"  - {'✓ verified' if st['ok'] else '✗ NOT verified'}"))
    if st.get("note"):
        lines.append(("note", f"  - _{st['note']}_"))
    return lines


def _diagrams_markdown_block(diagrams: dict) -> str:
    """One heading + embedded image per (material, diagram) -- mirrors
    :func:`_summary_markdown_block`. Images are embedded as base64 data
    URIs (same trick :func:`render_html` uses for the per-report section
    figure) so this works whether the .md text is written straight to
    disk, piped through pandoc, or returned as a plain string -- no
    separate assets folder to keep track of."""
    import base64
    out = []
    for label, shots in diagrams.items():
        if not shots:
            continue
        out.append(f"## {label} -- diagrams\n")
        for caption, png in shots:
            b64 = base64.b64encode(png).decode("ascii")
            out.append(f"**{caption}**\n")
            out.append(f"![{caption}](data:image/png;base64,{b64})\n")
    return "\n".join(out)


def _summary_markdown_block(summary: dict) -> str:
    """One Markdown pipe-table per material -- ``{material_label: rows}``,
    *rows* from :func:`summary_rows` -- rendered before the narrative
    sections (General tab content, see design_report_panel.py). Column set
    follows each row-set's ``kind`` (see :func:`summary_rows`): reinforcement
    areas for concrete bars/areas, the utilization ratio for everything
    else."""
    out = []
    for label, rows in summary.items():
        if not rows:
            continue
        kind = rows[0].get("kind", "ratio")
        out.append(f"## {label} -- summary\n")
        if kind == "concrete_bar":
            out.append("| Id | Section | Combination | As,bot | As,top | Asw/s |")
            out.append("|---|---|---|---|---|---|")
            for r in rows:
                out.append(f"| {r['id']} | {r['section']} | {r['combination']} "
                           f"| {_fmt_area(r['As_bot'])} | {_fmt_area(r['As_top'])} "
                           f"| {_fmt_area(r['Asw_s'])} |")
        elif kind == "concrete_area":
            out.append("| Id | Section | Combination | Asx,bot | Asx,top | Asy,bot | Asy,top |")
            out.append("|---|---|---|---|---|---|---|")
            for r in rows:
                out.append(f"| {r['id']} | {r['section']} | {r['combination']} "
                           f"| {_fmt_area(r['Asx_bot'])} | {_fmt_area(r['Asx_top'])} "
                           f"| {_fmt_area(r['Asy_bot'])} | {_fmt_area(r['Asy_top'])} |")
        else:
            out.append("| Id | Section | Combination | Utilization | OK |")
            out.append("|---|---|---|---|---|")
            for r in rows:
                util = "--" if r["utilization"] is None else _fmt(r["utilization"])
                ok = "--" if r["ok"] is None else ("Yes" if r["ok"] else "**No**")
                out.append(f"| {r['id']} | {r['section']} | {r['combination']} "
                           f"| {util} | {ok} |")
        out.append("")
    return "\n".join(out)


def render_markdown(reports: dict, title: str = "Design report",
                    section_images: dict | None = None,
                    summary: dict | None = None,
                    diagrams: dict | None = None) -> str:
    """Render ``{id: report_dict}`` as Markdown.

    ``section_images`` optionally maps a report id to a (relative) image path to
    embed as the section drawing (see :func:`write_markdown_bundle`). Equations
    are kept as ``$$…$$`` LaTeX (math viewers / pandoc). ``summary`` and
    ``diagrams``, when given, are the General tab's content (see
    design_report_panel.py) -- ``summary`` is ``{material_label: rows}``
    (see :func:`summary_rows`), ``diagrams`` is ``{material_label:
    [(caption, png_bytes), ...]}`` -- both rendered before the narrative
    sections, tables first then diagrams.
    """
    section_images = section_images or {}
    out = [f"# {title}\n"]
    if summary:
        out.append(_summary_markdown_block(summary))
    if diagrams:
        out.append(_diagrams_markdown_block(diagrams))
    for _id, rep in reports.items():
        out.append(f"## {rep.get('title', _id)}\n")
        meta = rep.get("meta") or {}
        if meta:
            out.append(" · ".join(f"**{k}**: {v}"
                                  for k, v in _meta_display_items(meta)) + "\n")
        if _id in section_images:
            out.append(f"![section]({section_images[_id]})\n")
        for sec in rep.get("sections", []):
            if sec.get("title"):
                out.append(f"### {sec['title']}\n")
            for st in sec.get("steps", []):
                out.append(_step_head(st))
                # Emit the LaTeX as a display equation ($$…$$): renders in
                # math-aware Markdown viewers, and pandoc turns it into native
                # OMML when converting to .docx. Blank line first: a $$…$$
                # right after a paragraph line, with no blank in between, is
                # read by CommonMark/pandoc as a lazy continuation of that
                # same paragraph, not its own block.
                if st.get("latex"):
                    out.append("")
                    out.append(f"$$ {st['latex']} $$")
                # Same for the bullet list (expr/subst/ok/note): without a
                # blank line before its first "- " item, CommonMark/pandoc
                # treats it as a lazy continuation of the preceding paragraph
                # (or display equation) instead of a list of its own — the
                # item text gets swallowed/merged instead of rendering.
                lines = _step_lines(st)
                if lines:
                    out.append("")
                    out.extend(lines)
                # A blank line between steps: without it, CommonMark (and
                # pandoc) treats consecutive non-blank lines as ONE paragraph
                # / list, merging every step of a section into a single
                # unreadable blob instead of one block per step.
                out.append("")
    return "\n".join(out)


def write_markdown_bundle(reports: dict, md_path, title: str = "Design report",
                          summary: dict | None = None,
                          diagrams: dict | None = None):
    """Write a Markdown report *and* an assets folder next to it.

    Creates ``<stem>_assets/`` beside *md_path* with the section drawings as PNG
    files, and writes the ``.md`` referencing them — so the Markdown shows the
    drawings in any viewer (and equations render as LaTeX in math-aware ones /
    via pandoc). Returns the assets directory path (or ``None`` if no drawing).
    """
    import re
    from pathlib import Path
    md_path = Path(md_path)
    assets = md_path.with_name(md_path.stem + "_assets")
    images: dict = {}
    made = False
    for _id, rep in reports.items():
        geom = (rep.get("meta") or {}).get("section_geom")
        png = section_figure_png(geom) if geom else None
        if not png:
            continue
        assets.mkdir(exist_ok=True)
        made = True
        slug = re.sub(r"\W+", "_", str(_id)).strip("_") or "sec"
        (assets / f"{slug}_section.png").write_bytes(png)
        images[_id] = f"{assets.name}/{slug}_section.png"
    md_path.write_text(render_markdown(reports, title, section_images=images,
                                       summary=summary, diagrams=diagrams),
                       encoding="utf-8")
    return assets if made else None


# ── HTML (MathJax) ───────────────────────────────────────────────────────────

_HTML_HEAD = """<!doctype html><html><head><meta charset="utf-8">
<title>{title}</title>
<script>
window.MathJax = {{ tex: {{ inlineMath: [['\\\\(','\\\\)']] }} }};
</script>
<script src="https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-mml-chtml.js" async></script>
<style>
 body{{font-family:{css_family};margin:2rem;color:#1a1a1a;font-size:{font_size}pt;}}
 h1{{border-bottom:2px solid #888;}} h2{{margin-top:2rem;color:#1f4fd8;}}
 h3{{margin:1rem 0 .3rem;color:#444;}}
 .meta{{color:#555;font-size:.9em;margin:.3rem 0 1rem;}}
 .step{{margin:.15rem 0;}} .sym{{font-weight:600;}}
 .clause{{color:#888;font-size:.85em;}}
 .expr,.subst{{color:#333;margin-left:1.5rem;font-size:.95em;}}
 .ok{{color:#137333;margin-left:1.5rem;}} .bad{{color:#c5221f;margin-left:1.5rem;}}
 .note{{color:#777;font-style:italic;margin-left:1.5rem;}}
</style></head><body>
<h1>{title}</h1>
"""


def _summary_html_block(summary: dict, esc) -> str:
    """One HTML ``<table>`` per material -- mirrors :func:`_summary_markdown_block`,
    including its per-``kind`` column set."""
    parts = []
    for label, rows in summary.items():
        if not rows:
            continue
        kind = rows[0].get("kind", "ratio")
        parts.append(f"<h3>{esc(label)} \u2014 summary</h3>")
        parts.append('<table border="1" cellpadding="4" cellspacing="0" '
                     'style="border-collapse:collapse;margin-bottom:1rem;">')
        if kind == "concrete_bar":
            parts.append("<tr><th>Id</th><th>Section</th><th>Combination</th>"
                         "<th>As,bot</th><th>As,top</th><th>Asw/s</th></tr>")
            for r in rows:
                parts.append(
                    f"<tr><td>{esc(str(r['id']))}</td><td>{esc(str(r['section']))}</td>"
                    f"<td>{esc(str(r['combination']))}</td>"
                    f"<td>{esc(_fmt_area(r['As_bot']))}</td>"
                    f"<td>{esc(_fmt_area(r['As_top']))}</td>"
                    f"<td>{esc(_fmt_area(r['Asw_s']))}</td></tr>")
        elif kind == "concrete_area":
            parts.append("<tr><th>Id</th><th>Section</th><th>Combination</th>"
                         "<th>Asx,bot</th><th>Asx,top</th><th>Asy,bot</th>"
                         "<th>Asy,top</th></tr>")
            for r in rows:
                parts.append(
                    f"<tr><td>{esc(str(r['id']))}</td><td>{esc(str(r['section']))}</td>"
                    f"<td>{esc(str(r['combination']))}</td>"
                    f"<td>{esc(_fmt_area(r['Asx_bot']))}</td>"
                    f"<td>{esc(_fmt_area(r['Asx_top']))}</td>"
                    f"<td>{esc(_fmt_area(r['Asy_bot']))}</td>"
                    f"<td>{esc(_fmt_area(r['Asy_top']))}</td></tr>")
        else:
            parts.append("<tr><th>Id</th><th>Section</th><th>Combination</th>"
                         "<th>Utilization</th><th>OK</th></tr>")
            for r in rows:
                util = "\u2014" if r["utilization"] is None else esc(_fmt(r["utilization"]))
                ok_cls = "" if r["ok"] is None else (' class="ok"' if r["ok"] else ' class="bad"')
                ok_txt = "\u2014" if r["ok"] is None else ("Yes" if r["ok"] else "No")
                parts.append(
                    f"<tr><td>{esc(str(r['id']))}</td><td>{esc(str(r['section']))}</td>"
                    f"<td>{esc(str(r['combination']))}</td><td>{util}</td>"
                    f"<td{ok_cls}>{ok_txt}</td></tr>")
        parts.append("</table>")
    return "\n".join(parts)


def _diagrams_html_block(diagrams: dict, esc) -> str:
    """One heading + embedded ``<img>`` per (material, diagram) -- mirrors
    :func:`_summary_html_block`."""
    import base64
    parts = []
    for label, shots in diagrams.items():
        if not shots:
            continue
        parts.append(f"<h3>{esc(label)} \u2014 diagrams</h3>")
        for caption, png in shots:
            b64 = base64.b64encode(png).decode("ascii")
            parts.append(f'<div class="meta">{esc(caption)}</div>')
            parts.append(f'<img alt="{esc(caption)}" '
                         f'src="data:image/png;base64,{b64}" '
                         'style="max-width:100%;margin:.3rem 0 1rem;">')
    return "\n".join(parts)


def render_html(reports: dict, title: str = "Design report",
                font_size: int = DEFAULT_REPORT_FONT_SIZE,
                font_family: str = DEFAULT_REPORT_FONT_FAMILY,
                summary: dict | None = None,
                diagrams: dict | None = None) -> str:
    """Render ``{id: report_dict}`` as a self-contained HTML page (MathJax).

    *font_size* (pt) and *font_family* (a key into REPORT_FONT_CHOICES) set
    the page's body font via CSS; MathJax-rendered equations inherit the
    surrounding font stack automatically. ``summary``/``diagrams``, when
    given, are the General tab's content (see :func:`summary_rows` and
    design_report_panel.py), rendered before the narrative sections,
    tables first then diagrams.
    """
    esc = _html.escape
    _font = _report_font_choice(font_family)
    parts = [_HTML_HEAD.format(title=esc(title), css_family=_font["css_family"],
                               font_size=font_size)]
    if summary:
        parts.append(_summary_html_block(summary, esc))
    if diagrams:
        parts.append(_diagrams_html_block(diagrams, esc))
    # Schema-mismatch banner (non-fatal).
    mism = schema_mismatches(reports)
    if mism:
        parts.append(
            '<div class="bad">⚠ Some reports use a different schema version '
            f'(expected {EXPECTED_SCHEMA}): {esc(", ".join(map(str, mism)))}. '
            'The rendering may be incomplete.</div>')
    # Table of contents (when more than one report).
    if len(reports) > 1:
        parts.append("<h3>Contents</h3><ul>")
        for i, (rid, rep) in enumerate(reports.items()):
            parts.append(f'<li><a href="#r{i}">'
                         f'{esc(str(rep.get("title", rid)))}</a></li>')
        parts.append("</ul>")
    for i, (_id, rep) in enumerate(reports.items()):
        parts.append(f'<h2 id="r{i}">{esc(str(rep.get("title", _id)))}</h2>')
        meta = rep.get("meta") or {}
        if meta:
            parts.append('<div class="meta">'
                         + " · ".join(f"{esc(str(k))}: {esc(v)}"
                                      for k, v in _meta_display_items(meta))
                         + "</div>")
        png = section_figure_png(meta.get("section_geom")) if meta else None
        if png:
            import base64
            b64 = base64.b64encode(png).decode("ascii")
            parts.append(f'<img alt="section" src="data:image/png;base64,{b64}" '
                         'style="height:150px;margin:.3rem 0;">')
        for sec in rep.get("sections", []):
            if sec.get("title"):
                parts.append(f"<h3>{esc(sec['title'])}</h3>")
            for st in sec.get("steps", []):
                clause = (f'  <span class="clause">[{esc(st["clause"])}]</span>'
                          if st.get("clause") else "")
                # Same "M_Ed" -> subscript convention as Markdown/PDF/Word
                # (_texify), carried into the page via MathJax's own inline
                # delimiters (configured above as \( \), not $...$) so the
                # symbol/expr/subst render as real subscripts here too,
                # instead of showing the raw "V_pl,Rd,y" text.
                sym_tex = esc(_texify(st["symbol"]))
                parts.append(
                    f'<div class="step"><span class="sym">\\(\\mathbf{{{sym_tex}}}\\)</span> '
                    f'= {esc(_fmt(st["value"]))} {esc(st.get("unit",""))}{clause}</div>')
                if st.get("latex"):
                    parts.append(f'<div class="expr">\\({st["latex"]}\\)</div>')
                elif st.get("expr"):
                    expr_tex = esc(_texify(st["expr"]))
                    parts.append(f'<div class="expr">\\({expr_tex}\\)</div>')
                if st.get("subst"):
                    subst_tex = esc(_texify(st["subst"]))
                    parts.append(f'<div class="subst">= \\({subst_tex}\\)</div>')
                if st.get("ok") is not None:
                    cls, txt = ("ok", "✓ verified") if st["ok"] else ("bad", "✗ NOT verified")
                    parts.append(f'<div class="{cls}">{txt}</div>')
                if st.get("note"):
                    parts.append(f'<div class="note">{esc(st["note"])}</div>')
    parts.append("</body></html>")
    return "\n".join(parts)


# ── PDF ──────────────────────────────────────────────────────────────────────

def _pdf_text_line(line: str) -> str:
    """Strip our own Markdown-only decoration — bullet ``- ``, the ``_[…]_``
    / ``_…_`` italic wrapper — from a ``_step_head``/``_step_lines`` string,
    for plain matplotlib ``text()``. Any ``$…$``/``$\\mathbf{…}$`` math
    segments are left untouched: matplotlib's mathtext renders them directly
    (no LaTeX/pandoc needed), so the PDF gets the same real subscripts as the
    other formats — and, since it's built from the exact same helper
    functions, the same wording too."""
    s = line.lstrip()
    if s.startswith("- "):
        s = s[2:]
    s = _re.sub(r'_\[(.*?)\]_', r'[\1]', s)          # "_[clause]_" -> "[clause]"
    if s.startswith("_") and s.endswith("_") and len(s) > 1:
        s = s[1:-1]                                   # whole-line "_note_"
    return s


def render_pdf(reports: dict, path: str, title: str = "Design report",
               progress_cb=None, font_size: int = DEFAULT_REPORT_FONT_SIZE,
               font_family: str = DEFAULT_REPORT_FONT_FAMILY,
               summary: dict | None = None,
               diagrams: dict | None = None) -> None:
    """Compose an A4 PDF with matplotlib: text + **native, vector-typeset
    equations** (:func:`_mathtext_extent_in` for sizing + a direct
    ``Figure.text`` for the actual draw — no rasterisation) + the **section
    figure** (:func:`section_figure_png`, a genuine picture, kept as PNG).

    Self-contained — no LaTeX / pandoc. Every text line is built by the same
    ``_step_head``/``_step_lines`` helpers :func:`render_markdown` uses (via
    :func:`_pdf_text_line`, which strips the Markdown-only decoration but
    keeps the ``$…$`` math — matplotlib's mathtext renders that directly), so
    the PDF reads identically to the ``.md``/Word/HTML exports.

    Equations used to be typeset via :func:`latex_png` (a 150dpi PNG) and
    embedded with ``imshow`` — crisp on screen at 100%, visibly pixelated at
    any zoom or at print resolution, unlike the vector text/lines around it.
    Since this whole page is already a matplotlib ``Figure`` saved through
    ``PdfPages``, an equation can instead be drawn with ``Figure.text`` like
    any other line — mathtext then stays vector in the PDF, exactly as sharp
    as the surrounding text. :func:`latex_png` remains as the fallback for
    the rare expression mathtext's ``path`` layout can size but the plain
    ``Figure.text`` draw call rejects.

    *progress_cb*, if given, is called as ``progress_cb(done, total)`` once
    per report (``done`` from 1..``total``) -- the granularity a caller on a
    background thread needs to drive a progress bar without adding overhead
    per equation. Called with plain ints; safe to marshal to the GUI thread
    via a Qt signal.

    *font_size* is the base body-text size in points (headings/meta/
    equations scale off it, same ratios as the old hardcoded 8/9/11/14/7
    — passing the defaults reproduces them exactly). *font_family* is a
    key into REPORT_FONT_CHOICES ('sans'/'serif'/'stix'/'cm'); it drives
    both matplotlib's rcParams (font.family/mathtext.fontset, restored
    when this function returns) so plain text and equations use a
    matching typeface.
    """
    from io import BytesIO
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.image as mpimg
    from matplotlib.backends.backend_pdf import PdfPages
    from matplotlib.figure import Figure

    # Same ratios the old hardcoded sizes had at their implicit base of
    # 10pt (8/9/11/14/7) — passing the default font_size reproduces them
    # exactly; a different font_size scales all of them together.
    title_size = font_size + 4
    report_title_size = font_size + 1
    meta_size = font_size - 3
    subsec_size = font_size - 1
    body_size = font_size - 2
    eq_size = font_size - 1

    A4 = (8.27, 11.69)
    M = 0.06                         # page margin (figure fraction)
    st = {"fig": None, "y": 0.0}

    _font = _report_font_choice(font_family)
    _old_rc = {"font.family": matplotlib.rcParams["font.family"],
               "mathtext.fontset": matplotlib.rcParams["mathtext.fontset"]}
    matplotlib.rcParams["font.family"] = _font["mpl_family"]
    matplotlib.rcParams["mathtext.fontset"] = _font["mathtext_fontset"]

    try:
        with PdfPages(path) as pdf:
            def _flush():
                if st["fig"] is not None:
                    pdf.savefig(st["fig"])
            def _newpage():
                _flush()
                fig = Figure(figsize=A4)
                st["fig"], st["y"] = fig, 1.0 - M
            def _ensure(space):
                if st["fig"] is None or st["y"] - space < M:
                    _newpage()
            def _text(s, size=body_size, bold=False, indent=0.0, color="black"):
                _ensure(size / 72.0 / A4[1] * 2.0)
                st["fig"].text(M + indent, st["y"], s, fontsize=size, va="top",
                               ha="left", weight=("bold" if bold else "normal"),
                               color=color, wrap=True)
                st["y"] -= (size / 72.0 / A4[1]) * 2.0
            def _image(png, max_h_in=0.30):
                img = mpimg.imread(BytesIO(png))
                ih, iw = img.shape[0], img.shape[1]
                h = max_h_in / A4[1]
                w = min(h * (iw / ih), 1.0 - 2 * M)
                _ensure(h * 1.15)
                ax = st["fig"].add_axes([M, st["y"] - h, w, h])
                ax.imshow(img); ax.axis("off")
                st["y"] -= h * 1.2
            def _equation(latex, size=eq_size, indent=0.05):
                if not latex:
                    return
                norm = _normalize_mathtext(latex)
                try:
                    _w_in, h_in = _mathtext_extent_in(norm, size)
                    h = h_in / A4[1]
                    _ensure(h * 1.3)
                    st["fig"].text(M + indent, st["y"], f"${norm}$",
                                   fontsize=size, va="top", ha="left")
                    st["y"] -= h * 1.3
                except Exception:    # noqa: BLE001 — mathtext parse failure
                    eq = latex_png(latex)
                    if eq:
                        _image(eq, max_h_in=0.26)

            _text(title, size=title_size, bold=True)
            if summary:
                for label, rows in summary.items():
                    if not rows:
                        continue
                    _text(f"{label} \u2014 summary", size=report_title_size,
                          bold=True, color="#1f4fd8")
                    for r in rows:
                        util = "\u2014" if r["utilization"] is None else _fmt(r["utilization"])
                        ok = ("\u2014" if r["ok"] is None
                              else ("OK" if r["ok"] else "NOT OK"))
                        okcolor = ("#555" if r["ok"] is None
                                  else ("#137333" if r["ok"] else "#c5221f"))
                        _text(f"{r['id']}  \u2014  {r['section']}  "
                              f"({r['combination']})  \u2014  util {util}  "
                              f"\u2014  {ok}", size=body_size, indent=0.02,
                              color=okcolor)
            if diagrams:
                for label, shots in diagrams.items():
                    if not shots:
                        continue
                    _text(f"{label} \u2014 diagrams", size=report_title_size,
                          bold=True, color="#1f4fd8")
                    for caption, png in shots:
                        _text(caption, size=body_size, indent=0.02)
                        _image(png, max_h_in=3.2)
            _total = len(reports)
            for _i, (_id, rep) in enumerate(reports.items(), start=1):
                if progress_cb is not None:
                    progress_cb(_i, _total)
                _text(str(rep.get("title", _id)), size=report_title_size, bold=True, color="#1f4fd8")
                meta = rep.get("meta") or {}
                if meta:
                    _text(" · ".join(f"{k}: {v}" for k, v in _meta_display_items(meta)),
                          size=meta_size, color="#555")
                fig_png = section_figure_png(meta.get("section_geom")) if meta else None
                if fig_png:
                    _image(fig_png, max_h_in=1.4)
                for sec in rep.get("sections", []):
                    if sec.get("title"):
                        _text(sec["title"], size=subsec_size, bold=True, color="#444")
                    _line_color = {"expr": "#333", "subst": "#333",
                                  "ok": "#137333", "not_ok": "#c5221f",
                                  "note": "#777"}
                    for step in sec.get("steps", []):
                        _text(_pdf_text_line(_step_head(step)), size=body_size, indent=0.02)
                        if step.get("latex"):
                            _equation(step["latex"])
                        for kind, line in _step_lines_kv(step):
                            _text(_pdf_text_line(line), size=body_size, indent=0.05,
                                  color=_line_color[kind])
            _flush()
    finally:
        matplotlib.rcParams.update(_old_rc)


def pandoc_available() -> bool:
    """True when a ``pandoc`` binary is on PATH."""
    import shutil
    return shutil.which("pandoc") is not None


def render_docx_via_pandoc(reports: dict, path: str,
                           title: str = "Design report",
                           font_size: int = DEFAULT_REPORT_FONT_SIZE,
                           font_family: str = DEFAULT_REPORT_FONT_FAMILY,
                           summary: dict | None = None,
                           diagrams: dict | None = None) -> None:
    """Write Word by piping the LaTeX Markdown through **pandoc**, which turns
    the ``$$…$$`` equations into **native, editable OMML** — no MML2OMML.XSL /
    lxml needed. Requires a ``pandoc`` binary on PATH (raises otherwise).

    Note: this route carries the maths but not the embedded section figure
    (Markdown has no image bytes); use :func:`render_docx` for figures + image
    equations without an external tool.

    *font_size*/*font_family* are applied via a tiny, throwaway
    ``--reference-doc`` (built with ``python-docx``, the same
    ``REPORT_FONT_CHOICES``/``Normal``-style mechanism :func:`render_docx`
    uses) — pandoc still does all the actual Markdown→docx conversion
    (including the OMML equations); the reference doc only supplies the
    default font pandoc's own output styles inherit from.
    """
    import os
    import shutil
    import subprocess
    import tempfile
    if shutil.which("pandoc") is None:
        raise RuntimeError("pandoc not found on PATH")
    from docx import Document
    from docx.shared import Pt
    _font = _report_font_choice(font_family)
    md = render_markdown(reports, title, summary=summary, diagrams=diagrams)
    fd, mdpath = tempfile.mkstemp(suffix=".md")
    ref_fd, ref_path = tempfile.mkstemp(suffix=".docx")
    os.close(ref_fd)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(md)
        ref_doc = Document()
        ref_doc.styles["Normal"].font.size = Pt(font_size)
        ref_doc.styles["Normal"].font.name = _font["docx_name"]
        ref_doc.save(ref_path)
        subprocess.run(["pandoc", mdpath, "-o", path,
                       f"--reference-doc={ref_path}"], check=True)
    finally:
        for p in (mdpath, ref_path):
            try:
                os.remove(p)
            except OSError:
                pass


def render_html_via_pandoc(reports: dict, path: str,
                           title: str = "Design report",
                           font_size: int = DEFAULT_REPORT_FONT_SIZE,
                           font_family: str = DEFAULT_REPORT_FONT_FAMILY,
                           summary: dict | None = None,
                           diagrams: dict | None = None) -> None:
    """Write HTML by piping the LaTeX Markdown through **pandoc** (``--mathjax``,
    same CDN as :func:`render_html`'s own template) so the exported page comes
    straight from the same ``.md`` text instead of a second, hand-written HTML
    renderer that can drift out of sync with it. Requires a ``pandoc`` binary
    on PATH (raises otherwise).

    *font_size*/*font_family* are applied via ``--include-in-header``: a tiny
    throwaway ``<style>`` block (same ``css_family``/pt size
    :func:`render_html`'s own CSS uses) that pandoc inlines into the
    ``<head>`` of the ``--standalone`` page it produces, so the result stays
    a single self-contained file.
    """
    import os
    import shutil
    import subprocess
    import tempfile
    if shutil.which("pandoc") is None:
        raise RuntimeError("pandoc not found on PATH")
    _font = _report_font_choice(font_family)
    md = render_markdown(reports, title, summary=summary, diagrams=diagrams)
    fd, mdpath = tempfile.mkstemp(suffix=".md")
    hfd, header_path = tempfile.mkstemp(suffix=".html")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            f.write(md)
        with os.fdopen(hfd, "w", encoding="utf-8") as f:
            f.write(f"<style>body{{font-family:{_font['css_family']};"
                    f"font-size:{font_size}pt;}}</style>")
        subprocess.run(
            ["pandoc", mdpath, "-o", path, "--standalone",
             "--mathjax=https://cdn.jsdelivr.net/npm/mathjax@3/es5/tex-mml-chtml.js",
             "--include-in-header", header_path,
             "--metadata", f"title={title}"],
            check=True)
    finally:
        for p in (mdpath, header_path):
            try:
                os.remove(p)
            except OSError:
                pass


def export_docx(reports: dict, path: str, title: str = "Design report",
                font_size: int = DEFAULT_REPORT_FONT_SIZE,
                font_family: str = DEFAULT_REPORT_FONT_FAMILY,
                summary: dict | None = None,
                diagrams: dict | None = None) -> None:
    """Write a Word file at *path*: via **pandoc** (native OMML equations,
    text matching the ``.md`` byte for byte) when a ``pandoc`` binary is on
    PATH, else the built-in :func:`render_docx` (no external tool, LaTeX
    typeset as images + Word-native subscript runs).

    *font_size*/*font_family* are honoured on both paths — on the pandoc
    path via a throwaway ``--reference-doc`` (see
    :func:`render_docx_via_pandoc`), on the built-in path via
    :func:`render_docx`'s own ``Normal`` style. *summary*/*diagrams*, when
    given, are the General tab's content (see :func:`summary_rows`),
    rendered before the narrative, tables then diagrams."""
    if pandoc_available():
        render_docx_via_pandoc(reports, path, title,
                               font_size=font_size, font_family=font_family,
                               summary=summary, diagrams=diagrams)
    else:
        render_docx(reports, path, title, font_size=font_size,
                   font_family=font_family, summary=summary,
                   diagrams=diagrams)


def export_html(reports: dict, path: str, title: str = "Design report",
                font_size: int = DEFAULT_REPORT_FONT_SIZE,
                font_family: str = DEFAULT_REPORT_FONT_FAMILY,
                summary: dict | None = None,
                diagrams: dict | None = None) -> None:
    """Write an HTML file at *path*: via **pandoc** when available (same
    ``.md`` text, same MathJax CDN), else the built-in :func:`render_html`.

    *font_size*/*font_family* are honoured on both paths — on the pandoc
    path via a throwaway ``--include-in-header`` stylesheet (see
    :func:`render_html_via_pandoc`), on the built-in path via
    :func:`render_html`'s own CSS. *summary*/*diagrams*, when given, are
    the General tab's content (see :func:`summary_rows`), rendered before
    the narrative, tables then diagrams."""
    if pandoc_available():
        render_html_via_pandoc(reports, path, title,
                               font_size=font_size, font_family=font_family,
                               summary=summary, diagrams=diagrams)
    else:
        from pathlib import Path
        Path(path).write_text(
            render_html(reports, title, font_size=font_size,
                       font_family=font_family, summary=summary,
                       diagrams=diagrams),
            encoding="utf-8")


def _add_texified_run(paragraph, text: str, bold: bool = False, italic: bool = False):
    """Add *text* to *paragraph* as one or more runs, turning the "M_Ed" /
    "f_yk" convention into a **real Word subscript run** (``font.subscript``)
    instead of a literal underscore — no external tool (pandoc/LaTeX) needed,
    and it matches what :func:`render_markdown` shows (same ``_texify``
    convention, just applied at the run level instead of as ``$…$``)."""
    if not text:
        return
    pos = 0
    for m in _SUBSCRIPT_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.bold = bold; r.italic = italic
        r = paragraph.add_run(m.group(1))
        r.bold = bold; r.italic = italic
        r = paragraph.add_run(m.group(2))
        r.bold = bold; r.italic = italic
        r.font.subscript = True
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.bold = bold; r.italic = italic


# Matches, in priority order: bold math ``$\mathbf{…}$``, plain inline math
# ``$…$``, Markdown bold ``**…**``, Markdown italic ``_…_``. Used to turn one
# of *our own* Markdown lines (from ``_step_head``/``_step_lines`` — a small,
# fully-controlled subset, not arbitrary Markdown) into native docx runs.
_INLINE_MD_RE = _re.compile(
    r'\$\\mathbf\{(.+?)\}\$|\$(.+?)\$|\*\*(.+?)\*\*|_(.+?)_')


def _md_line_to_docx(paragraph, line: str) -> None:
    """Add *line* (one of our own Markdown lines) to *paragraph* as native
    Word runs: ``$\\mathbf{…}$``/``$…$`` become a bold/plain subscript run via
    ``_add_texified_run`` (same "M_Ed" convention as the .md), ``**…**``/``_…_``
    become bold/italic runs — so the docx text is the *same string* as the
    Markdown, just carried by Word formatting instead of Markdown syntax."""
    pos = 0
    for m in _INLINE_MD_RE.finditer(line):
        if m.start() > pos:
            paragraph.add_run(line[pos:m.start()])
        if m.group(1) is not None:
            _add_texified_run(paragraph, m.group(1), bold=True)
        elif m.group(2) is not None:
            _add_texified_run(paragraph, m.group(2))
        elif m.group(3) is not None:
            paragraph.add_run(m.group(3)).bold = True
        elif m.group(4) is not None:
            paragraph.add_run(m.group(4)).italic = True
        pos = m.end()
    if pos < len(line):
        paragraph.add_run(line[pos:])


def render_docx(reports: dict, path: str, title: str = "Design report",
                font_size: int = DEFAULT_REPORT_FONT_SIZE,
                font_family: str = DEFAULT_REPORT_FONT_FAMILY,
                summary: dict | None = None,
                diagrams: dict | None = None) -> None:
    """Write ``{id: report_dict}`` to a Word document at *path*.

    Native ``python-docx`` (no pandoc needed). Every line of text is built by
    the exact same ``_step_head``/``_step_lines`` helpers :func:`render_markdown`
    uses, then carried into the docx via :func:`_md_line_to_docx` (Markdown
    syntax → native Word runs/subscripts) — so this export reads identically
    to the ``.md``/pandoc paths, just without pandoc's OMML equations. A
    LaTeX-typeset image (:func:`latex_png`) is added *in addition* to the
    ``latex`` display-equation line when present and matplotlib can parse it;
    it never replaces the text, so nothing is silently dropped if it fails.

    *font_size* (pt) and *font_family* (a key into REPORT_FONT_CHOICES) set
    the document's ``Normal`` style, so every default paragraph, heading and
    table inherits them.
    """
    from io import BytesIO
    from docx import Document
    from docx.shared import Inches, Pt

    _font = _report_font_choice(font_family)
    doc = Document()
    # Word's own default (Normal style, usually 11pt Calibri) reads noticeably
    # larger than the PDF/HTML exports of the same report — match them so all
    # formats read consistently regardless of which one is opened.
    doc.styles["Normal"].font.size = Pt(font_size)
    doc.styles["Normal"].font.name = _font["docx_name"]
    doc.add_heading(title, level=0)
    # Page breaks after the summary tables, after the diagrams, and between
    # each material's narrative section — so Word doesn't run a table's last
    # row into the first diagram, or one material's steps into the next
    # material's heading, mid-page. Tracked rather than unconditional so an
    # empty summary/diagrams dict (nothing was checked) doesn't leave a
    # blank page at the front of the document.
    wrote_summary = False
    wrote_diagrams = False
    if summary:
        for label, rows in summary.items():
            if not rows:
                continue
            wrote_summary = True
            kind = rows[0].get("kind", "ratio")
            doc.add_heading(f"{label} \u2014 summary", level=1)
            if kind == "concrete_bar":
                cols = ("Id", "Section", "Combination", "As,bot", "As,top", "Asw/s")
                keys = ("As_bot", "As_top", "Asw_s")
            elif kind == "concrete_area":
                cols = ("Id", "Section", "Combination", "Asx,bot", "Asx,top",
                        "Asy,bot", "Asy,top")
                keys = ("Asx_bot", "Asx_top", "Asy_bot", "Asy_top")
            else:
                cols = ("Id", "Section", "Combination", "Utilization", "OK")
                keys = None
            tbl = doc.add_table(rows=1, cols=len(cols))
            tbl.style = "Light Grid Accent 1"
            hdr = tbl.rows[0].cells
            for c, h in enumerate(cols):
                hdr[c].text = h
            for r in rows:
                cells = tbl.add_row().cells
                cells[0].text = str(r["id"])
                cells[1].text = str(r["section"])
                cells[2].text = str(r["combination"])
                if keys is not None:
                    for c, k in enumerate(keys, start=3):
                        cells[c].text = _fmt_area(r[k])
                else:
                    cells[3].text = ("\u2014" if r["utilization"] is None
                                     else _fmt(r["utilization"]))
                    cells[4].text = ("\u2014" if r["ok"] is None
                                     else ("Yes" if r["ok"] else "No"))
    if wrote_summary and reports:
        # Only worth a break if there is more content (diagrams or the
        # narrative) to push onto the next page.
        doc.add_page_break()
    if diagrams:
        for label, shots in diagrams.items():
            if not shots:
                continue
            wrote_diagrams = True
            doc.add_heading(f"{label} \u2014 diagrams", level=1)
            for caption, png in shots:
                doc.add_paragraph(caption)
                doc.add_picture(BytesIO(png), width=Inches(6.0))
    if wrote_diagrams and reports:
        doc.add_page_break()
    # Reports are keyed "material::rid" by the Design report panel
    # (design_report_panel.py's ``_dr_add``) -- every OTHER caller (tests,
    # the pandoc path's own markdown builder) uses plain ids with no "::",
    # which simply means no material grouping is detected and no mid-body
    # breaks are inserted, same as before this change.
    prev_material = None
    for _id, rep in reports.items():
        material = _id.split("::", 1)[0] if "::" in _id else None
        if material is not None and prev_material is not None and material != prev_material:
            doc.add_page_break()
        prev_material = material
        doc.add_heading(str(rep.get("title", _id)), level=1)
        meta = rep.get("meta") or {}
        if meta:
            doc.add_paragraph(" · ".join(f"{k}: {v}"
                                         for k, v in _meta_display_items(meta)))
        fig = section_figure_png(meta.get("section_geom")) if meta else None
        if fig:
            doc.add_picture(BytesIO(fig), height=Inches(1.6))
        for sec in rep.get("sections", []):
            if sec.get("title"):
                doc.add_heading(sec["title"], level=2)
            for st in sec.get("steps", []):
                p = doc.add_paragraph()
                _md_line_to_docx(p, _step_head(st))
                # A LaTeX display equation gets an extra typeset image (no
                # Markdown equivalent for an image, so it's not part of
                # _step_lines) — additive, never a replacement for the text.
                if st.get("latex"):
                    eq = latex_png(st["latex"])
                    if eq:
                        doc.add_picture(BytesIO(eq), height=Inches(0.30))
                for line in _step_lines(st):
                    pl = doc.add_paragraph()
                    pl.paragraph_format.left_indent = Inches(0.25)
                    # Strip the leading "- " (and "- = " for subst) Markdown
                    # bullet marker — the indent already conveys nesting.
                    text = line.lstrip()
                    if text.startswith("- "):
                        text = text[2:]
                    _md_line_to_docx(pl, text)
    doc.save(path)
